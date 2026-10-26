"""Convertisseur PPTX (python-pptx) / PPTX converter."""

from __future__ import annotations

import dataclasses
import io

import pytest
from convert_samples import AUTHOR, EMAIL, IBAN, NIR, PHONE, PHONE_INTL, make_pptx

from gdprguardian.convert import to_markdown
from gdprguardian.documents import DocumentOptions, DocumentTooLargeError

LOCAL = DocumentOptions(sandbox=False)


@pytest.fixture(scope="module")
def pptx() -> bytes:
    return make_pptx()


def convert(data: bytes, **options: object):
    return to_markdown(data, "deck.pptx", dataclasses.replace(LOCAL, **options))


def test_slides_titles_and_bullets(pptx: bytes) -> None:
    result = convert(pptx)
    assert result.kind == "pptx" and result.pages == 3
    assert "## Diapositive 1 : Présentation RH" in result.markdown
    assert f"- Point A\n- Contact {EMAIL}\n  - Détail niveau deux" in result.markdown
    assert "## Diapositive 3 : Données" in result.markdown


def test_tables_groups_charts_and_pictures(pptx: bytes) -> None:
    markdown = convert(pptx).markdown
    assert f"| Nom\\|Prénom | Téléphone |\n| --- | --- |\n| Dupont | {PHONE} |" in markdown
    assert "DANS-GROUPE" in markdown
    assert "| Catégorie | Ventes |" in markdown and "| T2 | 12.5 |" in markdown
    assert "[IMAGE: Logo ACME]" in markdown


def test_hidden_slides_follow_include_hidden(pptx: bytes) -> None:
    default = convert(pptx)
    assert "SECRET-CACHE" not in default.markdown and default.dropped["hidden_slides"] == 1
    shown = convert(pptx, include_hidden=True)
    assert "## Diapositive 2 (masquée) : Diapo cachée" in shown.markdown
    assert IBAN in shown.markdown and "hidden_slides" not in shown.dropped


def test_hidden_shapes_follow_include_hidden() -> None:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for title, hidden_title in (("Titre", False), ("TITRE-CACHE", True)):
        slide = presentation.slides.add_slide(presentation.slide_layouts[5])
        slide.shapes.title.text = title
        if hidden_title:
            slide.shapes.title._element.nvSpPr.cNvPr.set("hidden", "1")
    for text, hidden in (("FORME-VISIBLE", None), ("FORME-CACHEE", "1")):
        box = presentation.slides[0].shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1))
        box.text_frame.text = text
        if hidden:
            box._element.nvSpPr.cNvPr.set("hidden", hidden)
    buffer = io.BytesIO()
    presentation.save(buffer)
    default = convert(buffer.getvalue())
    assert "## Diapositive 1 : Titre" in default.markdown and "FORME-VISIBLE" in default.markdown
    assert "CACHE" not in default.markdown and "## Diapositive 2\n" in default.markdown + "\n"
    assert default.dropped["hidden_content"] == 2
    shown = convert(buffer.getvalue(), include_hidden=True).markdown
    assert "FORME-CACHEE" in shown and "## Diapositive 2 : TITRE-CACHE" in shown


def test_speaker_notes_follow_option(pptx: bytes) -> None:
    assert "### Notes\n\nNotes privées : NIR " + NIR in convert(pptx).markdown
    without = convert(pptx, include_speaker_notes=False)
    assert NIR not in without.markdown and without.dropped["speaker_notes"] == 1


def test_comments_are_opt_in_and_anonymous(pptx: bytes) -> None:
    default = convert(pptx)
    assert PHONE_INTL not in default.markdown and default.dropped["comments"] == 1
    shown = convert(pptx, include_comments=True)
    assert "### Commentaires\n\n- Commentaire PPTX : appeler le " + PHONE_INTL in shown.markdown


def test_names_authors_and_metadata_never_output(pptx: bytes) -> None:
    shown = convert(pptx, include_comments=True, include_hidden=True)
    assert AUTHOR not in shown.markdown  # propriétés, auteurs de commentaires, nom de forme / props, authors, names
    assert shown.dropped["metadata"] >= 1


def test_limits(pptx: bytes) -> None:
    with pytest.raises(DocumentTooLargeError, match="diapositives"):
        convert(pptx, max_pages=2)
    with pytest.raises(DocumentTooLargeError, match="cellules"):
        convert(pptx, max_cells=3)
