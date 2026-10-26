"""Documents volumineux pour mesurer la vitesse de conversion / large documents to measure conversion speed."""

from __future__ import annotations

import io

from convert_samples import EMAIL, IBAN, NIR, PHONE, RTF, line, make_odf, make_pdf, odf_content

SENTENCE = (f"Le salarié Jean Dupont (NIR {NIR}, IBAN {IBAN}, courriel {EMAIL}, tél. {PHONE}) a signé l'avenant "
            "relatif à sa rémunération et à ses congés.")


def big_pdf(pages: int = 100) -> bytes:
    content = []
    for number in range(pages):
        items = [line(f"Chapitre {number + 1}", 750, 16, bold=True)]
        y = 720.0
        for row in range(45):
            items.append(line(f"{row:02d} Ligne de texte courant numéro {row} de la page {number + 1}, avec un "
                              "contenu suffisamment long pour remplir la colonne.", y))
            y -= 14
        content.append(items)
    return make_pdf(content)


def big_docx(paragraphs: int = 600, tables: int = 20) -> bytes:
    import docx

    document = docx.Document()
    for index in range(paragraphs):
        if index % 30 == 0:
            document.add_heading(f"Section {index // 30 + 1}", 1)
        document.add_paragraph(SENTENCE)
    for _ in range(tables):
        table = document.add_table(rows=10, cols=5)
        for row in table.rows:
            for cell in row.cells:
                cell.text = "Dupont"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def big_pptx(slides: int = 50) -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    presentation = Presentation()
    for index in range(slides):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = f"Diapositive {index + 1}"
        frame = slide.placeholders[1].text_frame
        frame.text = SENTENCE
        for _ in range(4):
            frame.add_paragraph().text = SENTENCE
        table = slide.shapes.add_table(4, 4, Inches(1), Inches(5), Inches(6), Inches(1)).table
        for row in table.rows:
            for cell in row.cells:
                cell.text = "Dupont"
        slide.notes_slide.notes_text_frame.text = SENTENCE
    buffer = io.BytesIO()
    presentation.save(buffer)
    return buffer.getvalue()


def big_html(paragraphs: int = 3000) -> bytes:
    rows = "".join(f"<tr><td>{i}</td><td>Jean Dupont</td><td>{EMAIL}</td></tr>" for i in range(500))
    body = "".join(f"<h2>Titre {i}</h2><p>{SENTENCE} <b>gras</b> <a href='https://x.example/{i}'>lien</a></p>"
                   if i % 20 == 0 else f"<p>{SENTENCE}</p>" for i in range(paragraphs))
    return f"<html><head><title>T</title></head><body>{body}<table>{rows}</table></body></html>".encode()


def big_rtf(paragraphs: int = 40_000) -> bytes:
    paragraph = r"\pard\plain Salari\'e9 Jean Dupont, NIR " + NIR + r" {\b gras} fin\par" + "\n"
    return (RTF[:-1] + paragraph * paragraphs + "}").encode("latin-1")


def big_odt(paragraphs: int = 3000) -> bytes:
    body = "".join(f'<text:h text:outline-level="1">Titre {i}</text:h>' if i % 50 == 0
                   else f"<text:p>{SENTENCE}</text:p>" for i in range(paragraphs))
    return make_odf(odf_content(body))


def big_text(size: int = 5 << 20) -> bytes:
    return ((SENTENCE + "\n") * (size // len(SENTENCE))).encode("cp1252")


DOCUMENTS = {"pdf": big_pdf, "docx": big_docx, "pptx": big_pptx, "html": big_html, "rtf": big_rtf, "odt": big_odt,
             "txt": big_text}
