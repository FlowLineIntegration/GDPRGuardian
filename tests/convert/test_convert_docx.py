"""Convertisseur DOCX (mammoth + XML direct) / DOCX converter."""

from __future__ import annotations

import dataclasses
import io

import mammoth
import pytest
from convert_samples import AUTHOR, EMAIL, IBAN, NIR, PHONE, PHONE_INTL, make_docx, rewrite_zip, zip_bytes
from mammoth.docx import office_xml, xmlparser

from gdprguardian.convert import _docx, to_markdown
from gdprguardian.convert._docx import fast_parsing, parse_for_mammoth
from gdprguardian.documents import DocumentError, DocumentOptions, DocumentTooLargeError

LOCAL = DocumentOptions(sandbox=False)
DTD_DOCUMENT = (b'<?xml version="1.0"?><!DOCTYPE d [<!ENTITY nom "Bonjour">]><w:document xmlns:w="http://schemas.'
                b'openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>&nom;</w:t></w:r></w:p>'
                b"</w:body></w:document>")


@pytest.fixture(scope="module")
def docx() -> bytes:
    return make_docx()


@pytest.fixture(scope="module")
def result(docx: bytes):
    return to_markdown(docx, "contrat.docx", LOCAL)


def test_body_headings_and_tables(result) -> None:
    lines = result.markdown.splitlines()
    assert lines[0] == "# Contrat de travail"
    assert f"Salarié : Jean Dupont, NIR **{NIR}**" in lines
    assert "| Nom\\|Prénom | IBAN |" in lines  # « | » échappé dans les cellules / escaped in cells
    assert f"| Dupont\\|Jean | {IBAN} |" in lines


def test_headers_and_footers_appear_once(result) -> None:
    markdown = result.markdown
    assert "## En-têtes et pieds de page" in markdown
    assert markdown.count("Confidentiel - contact " + EMAIL) == 1
    assert "Pied de page - tél. " + PHONE in markdown
    assert "Première page - courriel " + EMAIL in markdown


def test_tracked_changes_are_accepted(result) -> None:
    assert "Version INSERE" in result.markdown
    assert "SUPPRIME-SECRET" not in result.markdown
    assert result.dropped["tracked_deletions"] == 1


def test_comments_are_opt_in_and_anonymous(docx: bytes, result) -> None:
    assert PHONE_INTL not in result.markdown
    assert result.dropped["comments"] == 1
    shown = to_markdown(docx, "c.docx", dataclasses.replace(LOCAL, include_comments=True))
    assert "## Commentaires\n\n- Rappeler au " + PHONE_INTL in shown.markdown
    assert "ZQ" not in shown.markdown  # initiales de l'auteur / author initials
    assert "comments" not in shown.dropped


def test_metadata_never_output(docx: bytes) -> None:
    shown = to_markdown(docx, "c.docx", dataclasses.replace(LOCAL, include_comments=True, include_hidden=True))
    assert AUTHOR not in shown.markdown
    assert shown.dropped["metadata"] >= 4


def test_images_become_placeholders_without_data(result) -> None:
    assert "[IMAGE: Photo du badge]" in result.markdown
    assert "base64" not in result.markdown and "data:" not in result.markdown
    assert result.dropped["images"] == 1


def test_text_boxes_and_simple_fields_skipped_by_mammoth_are_recovered(result) -> None:
    assert "## Zones de texte" in result.markdown
    assert "Zone de texte BOITE-TEXTE" in result.markdown
    assert "CHAMP-SIMPLE" in result.markdown


def test_fast_parser_builds_the_same_tree_as_minidom(docx: bytes, monkeypatch: pytest.MonkeyPatch) -> None:
    """Le remplaçant defusedxml de minidom ne change pas le rendu de mammoth / same mammoth output."""
    assert office_xml.parse_xml is parse_for_mammoth
    calls: list[int] = []
    monkeypatch.setattr(_docx, "_MINIDOM_PARSE", lambda *args: calls.append(1) or xmlparser.parse_xml(*args))
    with fast_parsing(None):
        fast = mammoth.convert_to_html(io.BytesIO(docx)).value
    assert not calls  # vraiment l'analyseur defusedxml / really the defusedxml parser
    slow = mammoth.convert_to_html(io.BytesIO(docx)).value
    assert calls and fast == slow


def test_mammoth_outside_our_conversions_keeps_its_own_parser(docx: bytes) -> None:
    # Importer gdprguardian ne change pas les appels à mammoth de l'application hôte (DTD comprises).
    # Importing gdprguardian leaves the host application's own mammoth calls unchanged.
    dtd = rewrite_zip(docx, {"word/document.xml": DTD_DOCUMENT})
    assert "Bonjour" in mammoth.convert_to_html(io.BytesIO(dtd)).value
    with pytest.raises(DocumentError):
        to_markdown(dtd, "c.docx", LOCAL)


def _hidden_docx() -> bytes:
    import docx as python_docx

    document = python_docx.Document()
    paragraph = document.add_paragraph("Visible ")
    paragraph.add_run("CACHE-CORPS").font.hidden = True
    paragraph.add_run(" suite")
    shown = paragraph.add_run(" MONTRE")
    shown.font.hidden = False  # <w:vanish w:val="0"/> : visible
    header = document.sections[0].header.paragraphs[0]
    header.text = "En-tête "
    header.add_run("CACHE-ENTETE").font.hidden = True
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_hidden_text_follows_include_hidden() -> None:
    data = _hidden_docx()
    default = to_markdown(data, "c.docx", LOCAL)
    assert "Visible suite MONTRE" in default.markdown and "- En-tête" in default.markdown
    assert "CACHE" not in default.markdown
    assert default.dropped["hidden_text"] == 2
    shown = to_markdown(data, "c.docx", dataclasses.replace(LOCAL, include_hidden=True))
    assert "Visible CACHE-CORPS suite MONTRE" in shown.markdown and "En-tête CACHE-ENTETE" in shown.markdown
    assert "hidden_text" not in shown.dropped


def test_fast_parser_drops_hidden_runs_only_when_asked() -> None:
    data = _hidden_docx()
    with fast_parsing(None):
        assert "CACHE-CORPS" in mammoth.convert_to_html(io.BytesIO(data)).value
    hidden = _docx._HiddenRuns()
    with fast_parsing(hidden):
        assert "CACHE-CORPS" not in mammoth.convert_to_html(io.BytesIO(data)).value
    assert hidden.count == 1
    assert "CACHE-CORPS" in mammoth.convert_to_html(io.BytesIO(data)).value


def test_damaged_docx_is_a_document_error(docx: bytes) -> None:
    broken = docx.replace(b"word/document.xml", b"word/documenX.xml")
    with pytest.raises(DocumentError):
        to_markdown(broken, "c.docx", LOCAL)


# --- limites comptées avant mammoth / limits counted before mammoth -----------------------------------------------

_W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_CONTENT_TYPES = (b'<?xml version="1.0" encoding="UTF-8"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/'
                  b'content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.'
                  b'relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName='
                  b'"/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.'
                  b'document.main+xml"/></Types>')
_PACKAGE_RELS = (b'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="http://schemas.openxmlformats.org/'
                 b'package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/'
                 b'officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')


def _body(body: str, prefix: str = "w") -> bytes:
    """DOCX minimal (ni commentaire ni note) au corps donné / minimal DOCX with the given body."""
    xml = (f'<?xml version="1.0" encoding="UTF-8"?><{prefix}:document xmlns:{prefix}="{_W_NS}"><{prefix}:body>'
           f"{body}</{prefix}:body></{prefix}:document>")
    return zip_bytes({"[Content_Types].xml": _CONTENT_TYPES, "_rels/.rels": _PACKAGE_RELS,
                      "word/document.xml": xml.encode()})


def _table(rows: int, cols: int, prefix: str = "w") -> str:
    cell = "<w:tc><w:p><w:r><w:t>Jean Dupont</w:t></w:r></w:p></w:tc>"
    table = "<w:tbl>" + ("<w:tr>" + cell * cols + "</w:tr>") * rows + "</w:tbl>"
    return table.replace("<w:", f"<{prefix}:").replace("</w:", f"</{prefix}:")


@pytest.fixture
def no_mammoth(monkeypatch: pytest.MonkeyPatch) -> None:
    """Refus attendu avant mammoth / refusal expected before mammoth."""
    def refuse(*_: object, **__: object) -> None:
        raise AssertionError("mammoth appelé / mammoth called")

    monkeypatch.setattr(_docx.mammoth, "convert_to_html", refuse)


def test_word_limits_are_published_values() -> None:
    assert _docx.MAX_PARAGRAPHS == 100_000 and _docx.MAX_ELEMENTS == 1_000_000  # deploy/README.md, site


def test_paragraphs_over_the_limit_are_refused_before_mammoth(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(_docx, "MAX_PARAGRAPHS", 12)
    # 4 lignes de 3 cellules : 12 paragraphes, admis (les w:pPr ne comptent pas) / 12 paragraphs, admitted.
    fits = _body(_table(4, 3).replace("<w:p>", "<w:p><w:pPr/>"))
    assert to_markdown(fits, "c.docx", LOCAL).markdown.count("Jean Dupont") == 12
    with monkeypatch.context() as patch:
        patch.setattr(_docx.mammoth, "convert_to_html", lambda *_, **__: pytest.fail("mammoth appelé / called"))
        # Un préfixe d'espace de noms quelconque compte aussi / any namespace prefix counts too.
        for prefix in ("w", "x"):
            with pytest.raises(DocumentTooLargeError, match="Word paragraphs > 12"):
                to_markdown(_body(_table(13, 1, prefix), prefix), "c.docx", LOCAL)


def test_notes_and_comments_count_too(docx: bytes, monkeypatch: pytest.MonkeyPatch) -> None:
    # Le document d'exemple a des commentaires : mammoth les lit, ils comptent / the sample has comments, they count.
    archive = _docx.SafeArchive(docx)
    main = _docx.main_part(archive, "word/document.xml")
    _docx.check_size(archive, main, b"")  # sous la limite / under the limit
    monkeypatch.setattr(_docx, "MAX_PARAGRAPHS", 0)
    with pytest.raises(DocumentTooLargeError, match="Word paragraphs > 0"):
        _docx.check_size(archive, main, b"")  # le document vide : les commentaires seuls / comments alone


def test_elements_over_the_limit_are_refused_before_mammoth(monkeypatch: pytest.MonkeyPatch,
                                                            no_mammoth: None) -> None:
    monkeypatch.setattr(_docx, "MAX_ELEMENTS", 500)
    # Un seul paragraphe de 300 passages (600 éléments) / a single paragraph of 300 runs (600 elements).
    runs = "<w:p>" + "<w:r><w:t>Jean </w:t></w:r>" * 300 + "</w:p>"
    with pytest.raises(DocumentTooLargeError, match="Word elements > 500"):
        to_markdown(_body(runs), "c.docx", LOCAL)


def test_converted_html_over_its_limit_names_the_word_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    from gdprguardian.convert import _html

    monkeypatch.setattr(_html, "MAX_TAGS", 40)
    with pytest.raises(DocumentTooLargeError, match="Word elements > 1000000"):
        to_markdown(_body(_table(10, 2)), "c.docx", LOCAL)
