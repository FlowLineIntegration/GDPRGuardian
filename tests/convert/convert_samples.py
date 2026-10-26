"""Fabriques de documents de test, construits en mémoire / in-memory test document builders.

Chaque document cache des identifiants dans un canal différent (corps, en-tête, commentaire, champ...).
Each document hides identifiers in a different channel (body, header, comment, form field...).
"""

from __future__ import annotations

import io
import os
import struct
import time
import zipfile
import zlib
from email.message import EmailMessage

NIR = "1 85 05 78 006 084 91"
IBAN = "FR76 3000 6000 0112 3456 7890 189"
EMAIL = "jean.dupont@exemple.fr"
PHONE = "06 12 34 56 78"
PHONE_INTL = "+33 6 98 76 54 32"
AUTHOR = "Zorglub Quixote"  # métadonnée indétectable : ne doit jamais sortir / undetectable metadata, never output
SECRETS = (NIR, IBAN, EMAIL, PHONE, PHONE_INTL)


def png_bytes() -> bytes:
    """PNG 1x1 valide sans Pillow / valid 1x1 PNG without Pillow."""
    def chunk(kind: bytes, payload: bytes) -> bytes:
        return struct.pack(">I", len(payload)) + kind + payload + struct.pack(">I", zlib.crc32(kind + payload))

    header = struct.pack(">IIBBBBB", 1, 1, 8, 2, 0, 0, 0)
    return (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", header) + chunk(b"IDAT", zlib.compress(b"\x00\xff\x00\x00"))
            + chunk(b"IEND", b""))


def zip_bytes(members: dict[str, bytes], compression: int = zipfile.ZIP_DEFLATED) -> bytes:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", compression) as archive:
        for name, data in members.items():
            archive.writestr(name, data)
    return buffer.getvalue()


def rewrite_zip(data: bytes, changes: dict[str, bytes]) -> bytes:
    """Remplace ou ajoute des entrées / replaces or adds members."""
    source = zipfile.ZipFile(io.BytesIO(data))
    members = {info.filename: source.read(info) for info in source.infolist()}
    members.update(changes)
    return zip_bytes(members)


# --- PDF ---------------------------------------------------------------------------------------------


def _pdf_string(text: str) -> bytes:
    raw = text.encode("cp1252")
    return b"(" + raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)") + b")"


def line(text: str, y: float, size: float = 10, bold: bool = False, x: float = 72, scale: float | None = None
         ) -> tuple[str, float, float, float, bool, float | None]:
    return text, size, x, y, bold, scale


def make_pdf(pages: list[list[tuple]], *, info: dict[str, str] | None = None, fields: dict[str, str] | None = None,
             annotations: list[tuple[str, str]] | None = None, scanned: frozenset[int] = frozenset(),
             attachments: dict[str, bytes] | None = None) -> bytes:
    """PDF minimal avec de vrais objets texte / minimal PDF with real text objects.

    `scale` dessine le texte en `/F1 1 Tf` avec une matrice (cas fréquent des PDF générés).
    `scale` draws text as `/F1 1 Tf` with a matrix (common in generated PDFs).
    """
    objects: list[bytes | None] = []

    def reserve() -> int:
        objects.append(None)
        return len(objects)

    def put(number: int, body: bytes) -> int:
        objects[number - 1] = body
        return number

    def add(body: bytes) -> int:
        return put(reserve(), body)

    def stream(dictionary: bytes, payload: bytes) -> bytes:
        return b"<< " + dictionary + b" /Length %d >>\nstream\n" % len(payload) + payload + b"\nendstream"

    catalog, pages_ref = reserve(), reserve()
    regular = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>")
    bold = add(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica-Bold /Encoding /WinAnsiEncoding >>")
    image = add(stream(b"/Type /XObject /Subtype /Image /Width 2 /Height 2 /ColorSpace /DeviceRGB "
                       b"/BitsPerComponent 8", bytes([200, 200, 200] * 4)))
    kids: list[int] = []
    field_refs: list[int] = []
    for index, items in enumerate(pages):
        page_ref = reserve()
        commands = []
        if index in scanned:
            commands.append(b"q 612 0 0 792 0 0 cm /Im1 Do Q")
        for text, size, x, y, is_bold, scale in items:
            font = b"/F2" if is_bold else b"/F1"
            if scale:
                commands.append(b"BT %g 0 0 %g %g %g Tm %s 1 Tf %s Tj ET" % (scale, scale, x, y, font,
                                                                            _pdf_string(text)))
            else:
                commands.append(b"BT %s %g Tf %g %g Td %s Tj ET" % (font, size, x, y, _pdf_string(text)))
        content = add(stream(b"", b"\n".join(commands)))
        annots: list[int] = []
        if index == 0:
            for name, value in (fields or {}).items():
                ref = add(b"<< /Type /Annot /Subtype /Widget /FT /Tx /T %s /V %s /Rect [72 100 300 120] /P %d 0 R "
                          b"/F 4 >>" % (_pdf_string(name), _pdf_string(value), page_ref))
                annots.append(ref)
                field_refs.append(ref)
            for subtype, contents in annotations or []:
                extra = b" /DA (/Helv 10 Tf 0 g)" if subtype == "FreeText" else b""
                annots.append(add(b"<< /Type /Annot /Subtype /%s /Contents %s /T %s /Rect [72 200 200 220]%s >>"
                                  % (subtype.encode(), _pdf_string(contents), _pdf_string(AUTHOR), extra)))
        annot_entry = b" /Annots [%s]" % b" ".join(b"%d 0 R" % ref for ref in annots) if annots else b""
        put(page_ref, b"<< /Type /Page /Parent %d 0 R /MediaBox [0 0 612 792] /Contents %d 0 R /Resources << /Font "
                      b"<< /F1 %d 0 R /F2 %d 0 R >> /XObject << /Im1 %d 0 R >> >>%s >>"
            % (pages_ref, content, regular, bold, image, annot_entry))
        kids.append(page_ref)
    put(pages_ref, b"<< /Type /Pages /Kids [%s] /Count %d >>" % (b" ".join(b"%d 0 R" % k for k in kids), len(kids)))
    extras = b""
    if field_refs:
        extras += b" /AcroForm << /Fields [%s] >>" % b" ".join(b"%d 0 R" % ref for ref in field_refs)
    if attachments:
        names = []
        for name, payload in attachments.items():
            embedded = add(stream(b"/Type /EmbeddedFile", payload))
            spec = add(b"<< /Type /Filespec /F %s /UF %s /EF << /F %d 0 R >> >>"
                       % (_pdf_string(name), _pdf_string(name), embedded))
            names.append(b"%s %d 0 R" % (_pdf_string(name), spec))
        extras += b" /Names << /EmbeddedFiles << /Names [%s] >> >>" % b" ".join(names)
    put(catalog, b"<< /Type /Catalog /Pages %d 0 R%s >>" % (pages_ref, extras))
    info_ref = add(b"<< %s >>" % b" ".join(b"/%s %s" % (key.encode(), _pdf_string(value))
                                           for key, value in info.items())) if info else None
    out = bytearray(b"%PDF-1.7\n%\xe2\xe3\xcf\xd3\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + (body or b"null") + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    trailer = b"/Size %d /Root %d 0 R" % (len(objects) + 1, catalog)
    if info_ref:
        trailer += b" /Info %d 0 R" % info_ref
    out += b"trailer\n<< %s >>\nstartxref\n%d\n%%%%EOF\n" % (trailer, xref)
    return bytes(out)


def encrypt_pdf(data: bytes, user_password: str, owner_password: str = "owner") -> bytes:
    import pypdf

    writer = pypdf.PdfWriter(clone_from=io.BytesIO(data))
    writer.encrypt(user_password, owner_password, algorithm="AES-128")
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def contract_pdf() -> bytes:
    """Contrat de deux pages : titres, paragraphe coupé, IBAN sur deux lignes, page numérisée, formulaire."""
    page_one = [
        line("Contrat de travail", 740, 20, bold=True),
        line("Article 1 - Parties", 700, 14, bold=True),
        line("Entre les soussignés Jean Dupont, numéro de sécurité sociale " + NIR + ", et la", 672),
        line("société ACME, il est convenu ce qui suit. Le salaire est versé sur le compte IBAN FR76 3000", 660),
        line("6000 0112 3456 7890 189 chaque mois.", 648),
        line("Adresse : 12 rue de la Paix", 620),
        line("75002 Paris", 608),
        line("Article 2 - Contact", 570, 14, bold=True),
        line("Courriel : " + EMAIL + ", tél. " + PHONE + ".", 548),
        line("Texte à l'échelle via une matrice de texte.", 500, scale=10),
    ]
    return make_pdf([page_one, []], info={"Author": AUTHOR, "Title": "Dossier " + AUTHOR},
                    fields={"Nom": "Jean Dupont", "NIR": NIR}, annotations=[("Text", "Rappeler au " + PHONE_INTL),
                                                                          ("FreeText", "Signé : Jean Dupont")],
                    scanned=frozenset({1}), attachments={"bulletin-dupont.pdf": b"%PDF-1.4 fake"})


# --- DOCX --------------------------------------------------------------------------------------------

_W_NS = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def make_docx() -> bytes:
    import docx
    from docx.oxml import parse_xml

    document = docx.Document()
    core = document.core_properties
    core.author = AUTHOR
    core.last_modified_by = AUTHOR
    core.title = "Dossier " + AUTHOR
    core.comments = "Note interne " + AUTHOR
    document.add_heading("Contrat de travail", 1)
    paragraph = document.add_paragraph("Salarié : Jean Dupont, NIR ")
    run = paragraph.add_run(NIR)
    run.bold = True
    document.add_comment(run, text="Rappeler au " + PHONE_INTL, author=AUTHOR, initials="ZQ")
    table = document.add_table(rows=2, cols=2)
    for (row, col), text in {(0, 0): "Nom|Prénom", (0, 1): "IBAN", (1, 0): "Dupont|Jean", (1, 1): IBAN}.items():
        table.cell(row, col).text = text
    section = document.sections[0]
    section.header.paragraphs[0].text = "Confidentiel - contact " + EMAIL
    section.footer.paragraphs[0].text = "Pied de page - tél. " + PHONE
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "Première page - courriel " + EMAIL
    picture = document.add_picture(io.BytesIO(png_bytes()))
    picture._inline.docPr.set("descr", "Photo du badge")
    changed = document.add_paragraph("Version ")
    changed._p.append(parse_xml(f'<w:ins {_W_NS} w:id="90" w:author="{AUTHOR}" w:date="2026-01-01T00:00:00Z">'
                                '<w:r><w:t>INSERE</w:t></w:r></w:ins>'))
    changed._p.append(parse_xml(f'<w:del {_W_NS} w:id="91" w:author="{AUTHOR}" w:date="2026-01-01T00:00:00Z">'
                                '<w:r><w:delText>SUPPRIME-SECRET</w:delText></w:r></w:del>'))
    # Zone de texte sans branche mc:Fallback, que mammoth ignore / text box without mc:Fallback, skipped by mammoth.
    box = (f'<w:p {_W_NS} xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"><w:r><mc:AlternateContent>'
           '<mc:Choice Requires="wps"><w:drawing><wps:wsp><wps:txbx><w:txbxContent><w:p><w:r><w:t>Zone de texte '
           'BOITE-TEXTE</w:t></w:r></w:p></w:txbxContent></wps:txbx></wps:wsp></w:drawing></mc:Choice>'
           '</mc:AlternateContent></w:r></w:p>')
    body = document.element.body
    body.insert(len(body) - 1, parse_xml(box))
    body.insert(len(body) - 1, parse_xml(f'<w:p {_W_NS}><w:fldSimple w:instr="AUTHOR"><w:r><w:t>CHAMP-SIMPLE</w:t>'
                                         '</w:r></w:fldSimple></w:p>'))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


# --- PPTX --------------------------------------------------------------------------------------------


def make_pptx() -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    presentation = Presentation()
    presentation.core_properties.author = AUTHOR
    first = presentation.slides.add_slide(presentation.slide_layouts[1])
    first.shapes.title.text = "Présentation RH"
    body = first.placeholders[1].text_frame
    body.text = "Point A"
    body.add_paragraph().text = "Contact " + EMAIL
    sub = body.add_paragraph()
    sub.text = "Détail niveau deux"
    sub.level = 1
    first.notes_slide.notes_text_frame.text = "Notes privées : NIR " + NIR
    hidden = presentation.slides.add_slide(presentation.slide_layouts[5])
    hidden.shapes.title.text = "Diapo cachée"
    hidden.shapes.add_textbox(Inches(1), Inches(2), Inches(4), Inches(1)).text_frame.text = "SECRET-CACHE " + IBAN
    hidden._element.set("show", "0")
    third = presentation.slides.add_slide(presentation.slide_layouts[5])
    third.shapes.title.text = "Données"
    table = third.shapes.add_table(2, 2, Inches(0.5), Inches(1.5), Inches(4), Inches(1)).table
    for (row, col), text in {(0, 0): "Nom|Prénom", (0, 1): "Téléphone", (1, 0): "Dupont", (1, 1): PHONE}.items():
        table.cell(row, col).text = text
    group = third.shapes.add_group_shape()
    group.shapes.add_textbox(0, 0, Inches(2), Inches(1)).text_frame.text = "DANS-GROUPE"
    picture = third.shapes.add_picture(io.BytesIO(png_bytes()), Inches(5), Inches(1))
    picture._element.nvPicPr.cNvPr.set("descr", "Logo ACME")
    picture._element.nvPicPr.cNvPr.set("name", "photo-" + AUTHOR)
    chart_data = CategoryChartData()
    chart_data.categories = ["T1", "T2"]
    chart_data.add_series("Ventes", (10.0, 12.5))
    third.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), chart_data)
    buffer = io.BytesIO()
    presentation.save(buffer)
    return _add_pptx_comment(buffer.getvalue())


def _add_pptx_comment(data: bytes) -> bytes:
    """python-pptx n'écrit pas les commentaires : ajout direct des parties / added as raw parts."""
    source = zipfile.ZipFile(io.BytesIO(data))
    p_ns = "http://schemas.openxmlformats.org/presentationml/2006/main"
    comment = (f'<?xml version="1.0" encoding="UTF-8" standalone="yes"?><p:cmLst xmlns:p="{p_ns}"><p:cm authorId="0" '
               'dt="2026-10-24T10:00:00.000" idx="1"><p:pos x="10" y="10"/><p:text>Commentaire PPTX : appeler le '
               f'{PHONE_INTL}</p:text></p:cm></p:cmLst>').encode()
    authors = (f'<?xml version="1.0" encoding="UTF-8"?><p:cmAuthorLst xmlns:p="{p_ns}"><p:cmAuthor id="0" '
               f'name="{AUTHOR}" initials="ZQ" lastIdx="1" clrIdx="0"/></p:cmAuthorLst>').encode()
    rels = source.read("ppt/slides/_rels/slide1.xml.rels").replace(
        b"</Relationships>", b'<Relationship Id="rId99" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
                             b'relationships/comments" Target="../comments/comment1.xml"/></Relationships>')
    types = source.read("[Content_Types].xml").replace(
        b"</Types>", b'<Override PartName="/ppt/comments/comment1.xml" ContentType="application/vnd.openxmlformats-'
                     b'officedocument.presentationml.comments+xml"/></Types>')
    return rewrite_zip(data, {"ppt/comments/comment1.xml": comment, "ppt/commentAuthors.xml": authors,
                              "ppt/slides/_rels/slide1.xml.rels": rels, "[Content_Types].xml": types})


# --- OpenDocument ------------------------------------------------------------------------------------

_ODF_NS = (
    'xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0" '
    'xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0" '
    'xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0" '
    'xmlns:draw="urn:oasis:names:tc:opendocument:xmlns:drawing:1.0" '
    'xmlns:style="urn:oasis:names:tc:opendocument:xmlns:style:1.0" '
    'xmlns:presentation="urn:oasis:names:tc:opendocument:xmlns:presentation:1.0" '
    'xmlns:svg="urn:oasis:names:tc:opendocument:xmlns:svg-compatible:1.0" '
    'xmlns:dc="http://purl.org/dc/elements/1.1/" '
    'xmlns:meta="urn:oasis:names:tc:opendocument:xmlns:meta:1.0" '
    'xmlns:xlink="http://www.w3.org/1999/xlink"'
)

ODT_BODY = f"""<text:tracked-changes><text:changed-region text:id="c1"><text:deletion>
<office:change-info><dc:creator>{AUTHOR}</dc:creator><dc:date>2026-01-01T00:00:00</dc:date></office:change-info>
<text:p>SUPPRIME-SECRET</text:p></text:deletion></text:changed-region></text:tracked-changes>
<text:h text:outline-level="1">Contrat de travail</text:h>
<text:p>Salarié : Jean<text:s text:c="2"/>Dupont, NIR {NIR}<text:note text:note-class="footnote" text:id="n1">
<text:note-citation>1</text:note-citation><text:note-body><text:p>Note : IBAN {IBAN}</text:p></text:note-body>
</text:note><text:change text:change-id="c1"/>.<office:annotation><dc:creator>{AUTHOR}</dc:creator>
<dc:date>2026-01-01T00:00:00</dc:date><text:p>Rappeler au {PHONE_INTL}</text:p></office:annotation></text:p>
<text:list><text:list-item><text:p>Premier point</text:p><text:list><text:list-item><text:p>Sous-point</text:p>
</text:list-item></text:list></text:list-item></text:list>
<table:table table:name="T1"><table:table-column table:number-columns-repeated="16384"/>
<table:table-header-rows><table:table-row><table:table-cell><text:p>Nom|Prénom</text:p></table:table-cell>
<table:table-cell><text:p>Courriel</text:p></table:table-cell></table:table-row></table:table-header-rows>
<table:table-row><table:table-cell><text:p>Dupont</text:p></table:table-cell><table:table-cell><text:p>{EMAIL}</text:p>
</table:table-cell><table:table-cell table:number-columns-repeated="16382"/></table:table-row>
<table:table-row table:number-rows-repeated="1048570"><table:table-cell table:number-columns-repeated="16384"/>
</table:table-row></table:table>
<text:p><draw:frame draw:name="photo-{AUTHOR}"><draw:image xlink:href="Pictures/a.png"/>
<svg:title>Photo badge</svg:title></draw:frame></text:p>"""


def odf_content(body: str, kind: str = "text", styles: str = "") -> bytes:
    return (f'<?xml version="1.0" encoding="UTF-8"?><office:document-content {_ODF_NS} office:version="1.3">'
            f"<office:automatic-styles>{styles}</office:automatic-styles>"
            f"<office:body><office:{kind}>{body}</office:{kind}></office:body></office:document-content>").encode()


def make_odf(content: bytes, *, mimetype: str = "application/vnd.oasis.opendocument.text",
             styles: bytes | None = None, manifest: bytes | None = None) -> bytes:
    styles = styles or (f'<?xml version="1.0" encoding="UTF-8"?><office:document-styles {_ODF_NS}>'
                        '<office:master-styles><style:master-page style:name="Standard"><style:header><text:p>'
                        f"En-tête : {EMAIL}</text:p></style:header><style:footer><text:p>Pied : tél. {PHONE}"
                        "</text:p></style:footer></style:master-page></office:master-styles>"
                        "</office:document-styles>").encode()
    meta = (f'<?xml version="1.0" encoding="UTF-8"?><office:document-meta {_ODF_NS}><office:meta>'
            f"<meta:initial-creator>{AUTHOR}</meta:initial-creator><dc:creator>{AUTHOR}</dc:creator>"
            "<dc:title>Titre secret</dc:title></office:meta></office:document-meta>").encode()
    manifest = manifest or (b'<?xml version="1.0" encoding="UTF-8"?><manifest:manifest xmlns:manifest="urn:oasis:'
                            b'names:tc:opendocument:xmlns:manifest:1.0"><manifest:file-entry manifest:full-path="/" '
                            b'manifest:media-type="' + mimetype.encode() + b'"/></manifest:manifest>')
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr(zipfile.ZipInfo("mimetype"), mimetype)  # stocké, en premier / stored, first
        for name, payload in (("content.xml", content), ("styles.xml", styles), ("meta.xml", meta),
                              ("META-INF/manifest.xml", manifest), ("Pictures/a.png", png_bytes())):
            archive.writestr(name, payload, zipfile.ZIP_DEFLATED)
    return buffer.getvalue()


def make_odt() -> bytes:
    return make_odf(odf_content(ODT_BODY))


def make_odp() -> bytes:
    hidden_style = ('<style:style style:name="dp2" style:family="drawing-page"><style:drawing-page-properties '
                    'presentation:visibility="hidden"/></style:style>')
    body = (f"<presentation:footer-decl presentation:name=\"ftr1\">Pied ODP {EMAIL}</presentation:footer-decl>"
            '<draw:page draw:name="page1" draw:style-name="dp1"><draw:frame presentation:class="title"><draw:text-box>'
            "<text:p>Bilan annuel</text:p></draw:text-box></draw:frame><draw:frame presentation:class=\"outline\">"
            f"<draw:text-box><text:list><text:list-item><text:p>Contact {PHONE}</text:p></text:list-item></text:list>"
            "</draw:text-box></draw:frame><presentation:notes><draw:page-thumbnail/><draw:frame "
            f'presentation:class="notes"><draw:text-box><text:p>Notes ODP {NIR}</text:p></draw:text-box></draw:frame>'
            '</presentation:notes></draw:page><draw:page draw:name="page2" draw:style-name="dp2"><draw:frame>'
            f"<draw:text-box><text:p>ODP-CACHE {IBAN}</text:p></draw:text-box></draw:frame></draw:page>")
    return make_odf(odf_content(body, "presentation", hidden_style),
                    mimetype="application/vnd.oasis.opendocument.presentation")


# --- RTF, HTML, EML ----------------------------------------------------------------------------------

def rtf_unicode(code: int) -> str:
    """Échappement RTF « \\uN » / RTF "\\uN" escape."""
    return "\\" + f"u{code}"


RTF = (r"{\rtf1\ansi\ansicpg1252\deff0{\fonttbl{\f0\fswiss\fcharset0 Arial;}{\f1\fnil\fcharset238 Arial CE;}}"
       r"{\info{\author " + AUTHOR + r"}{\title Dossier secret}{\operator " + AUTHOR + "}}"
       r"{\header \pard\plain Ent\'eate : " + EMAIL + r"\par}"
       r"\pard\plain\outlinelevel0 Contrat de travail\par"
       r"\pard\plain Salari\'e9 : Jean Dupont, NIR " + NIR + r"\par"
       r"\pard\plain Caf\'e9 " + rtf_unicode(8364) + r"? cr" + rtf_unicode(232)
       + r"\'e8me \f1 \'b3\'f3d\'9f\f0  {\deleted SUPPRIME-SECRET}"  # \fcharset238 : cp1250
       r"{\v TEXTE-MASQUE} {\*\atnid ZQ}{\*\atnauthor " + AUTHOR + r"}\chatn{\*\annotation\pard Rappeler au "
       + PHONE_INTL + r"\par}\par"
       r"{\pict\pngblip\bin6 }{\}\\x}\par"
       r"\pard\plain\ls1\ilvl0 Premier point\par"
       r"\pard\plain\intbl Nom|Pr\'e9nom\cell IBAN\cell\row"
       r"\pard\plain\intbl Dupont\cell " + IBAN + r"\cell\row"
       r"\pard\plain Fin {\field{\*\fldinst HYPERLINK " + '"mailto:' + EMAIL + '"' + r"}{\fldrslt courriel}}\par}")


HTML = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><meta name="author" content="{AUTHOR}">
<title>Dossier {AUTHOR}</title><style>p {{ color: red }}</style><script>var secret = "SCRIPT-SECRET";</script>
<link rel="stylesheet" href="https://tracker.example/style.css"></head><body>
<h1>Contrat</h1><p>Salarié : Jean Dupont, NIR {NIR}, <a href="mailto:{EMAIL}">écrire</a>, tél. {PHONE}.</p>
<!-- commentaire caché {AUTHOR} --><img src="https://tracker.example/pixel.gif" alt="Logo ACME">
<img src="data:image/png;base64,iVBORw0KGgo=" alt=""><iframe src="https://evil.example/"></iframe>
<table><tr><th>Nom|Prénom</th><th>IBAN</th></tr><tr><td>Dupont</td><td>{IBAN}</td></tr></table>
<div style="display: none">TEXTE-MASQUE</div><p hidden>AUSSI-MASQUE</p>
<ul><li>un</li><li>deux <b>gras</b></li></ul><p>jean_dupont_secours@exemple.fr</p></body></html>"""


def make_eml(*, html: bool = False, plain: bool = True) -> bytes:
    message = EmailMessage()
    message["From"] = f"Jean Dupont <{EMAIL}>"
    message["To"] = "Marie Martin <marie.martin@exemple.fr>"
    message["Cc"] = "rh@exemple.fr"
    message["Bcc"] = "cache@exemple.fr"
    message["Subject"] = "Dossier NIR " + NIR
    message["Date"] = "Sat, 24 Oct 2026 10:00:00 +0200"
    message["Message-ID"] = "<abc@mail.example>"
    message["Received"] = "from mail.example (192.168.1.42) by mx.example"
    message["X-Originating-IP"] = "[192.168.1.42]"
    if plain:
        message.set_content(f"Bonjour,\n\nmon IBAN est {IBAN}.\nTél. {PHONE}\n")
    if html:
        markup = f"<html><body><p>Version <b>HTML</b> : IBAN {IBAN}</p><script>x()</script></body></html>"
        if plain:
            message.add_alternative(markup, subtype="html")
        else:
            message.set_content(markup, subtype="html")
    message.add_attachment(b"%PDF-1.4 fake", maintype="application", subtype="pdf", filename="bulletin-dupont.pdf")
    return message.as_bytes()


# --- OLE (CFB) et MSG / OLE (CFB) and MSG --------------------------------------------------------------

_FREE, _END, _FAT = 0xFFFFFFFF, 0xFFFFFFFE, 0xFFFFFFFD


def make_cfb(tree: dict) -> bytes:
    """Fichier composé (CFB v3) minimal ; `tree` : {nom: octets | sous-arbre} / minimal compound file."""
    entries: list[dict] = [{"name": "Root Entry", "type": 5, "children": [], "data": b""}]

    def add(node: dict, parent: int) -> None:
        for name, value in node.items():
            entries.append({"name": name, "type": 1 if isinstance(value, dict) else 2, "children": [],
                            "data": b"" if isinstance(value, dict) else value})
            index = len(entries) - 1
            entries[parent]["children"].append(index)
            if isinstance(value, dict):
                add(value, index)

    add(tree, 0)
    for entry in entries:
        entry.update(left=_FREE, right=_FREE, child=_FREE, start=_END)

    def balance(ids: list[int]) -> int:
        if not ids:
            return _FREE
        middle = len(ids) // 2
        root = ids[middle]
        entries[root]["left"], entries[root]["right"] = balance(ids[:middle]), balance(ids[middle + 1:])
        return root

    for entry in entries:
        ordered = sorted(entry["children"], key=lambda i: (len(entries[i]["name"]), entries[i]["name"].upper()))
        entry["child"] = balance(ordered)
    # Petits flux dans le mini-flux (secteurs de 64 o) / small streams in the mini stream (64-byte sectors)
    mini = bytearray()
    mini_fat: list[int] = []
    big: list[dict] = []
    for entry in entries[1:]:
        data = entry["data"]
        if entry["type"] != 2 or not data:
            continue
        if len(data) < 4096:
            first = len(mini) // 64
            count = -(-len(data) // 64)
            mini_fat.extend([*range(first + 1, first + count), _END])
            mini += data + b"\0" * (count * 64 - len(data))
            entry["start"] = first
        else:
            big.append(entry)
    directory = -(-len(entries) * 128 // 512)
    mini_fat_sectors = -(-len(mini_fat) * 4 // 512)
    mini_sectors = -(-len(mini) // 512)
    big_sectors = [-(-len(entry["data"]) // 512) for entry in big]
    payload = directory + mini_fat_sectors + mini_sectors + sum(big_sectors)
    fat_sectors = 1
    while fat_sectors * 128 < fat_sectors + payload:
        fat_sectors += 1
    fat = [_FAT] * fat_sectors
    cursor = fat_sectors

    def chain(count: int) -> int:
        nonlocal cursor
        if count == 0:
            return _END
        start = cursor
        fat.extend([*range(cursor + 1, cursor + count), _END])
        cursor += count
        return start

    directory_start = chain(directory)
    mini_fat_start = chain(mini_fat_sectors)
    entries[0]["start"] = chain(mini_sectors)
    entries[0]["size"] = len(mini)
    for entry, count in zip(big, big_sectors, strict=True):
        entry["start"] = chain(count)
    fat += [_FREE] * (fat_sectors * 128 - len(fat))
    header = struct.pack("<8s16sHHHHH6sIIIIIIIII", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1", b"\0" * 16, 0x3E, 3, 0xFFFE,
                         9, 6, b"\0" * 6, 0, fat_sectors, directory_start, 0, 4096,
                         mini_fat_start if mini_fat else _END, mini_fat_sectors, _END, 0)
    header += struct.pack("<109I", *(list(range(fat_sectors)) + [_FREE] * (109 - fat_sectors)))

    def pad(blob: bytes) -> bytes:
        return blob + b"\0" * (-len(blob) % 512)

    dir_blob = bytearray()
    for entry in entries:
        name = entry["name"].encode("utf-16-le") + b"\0\0"
        size = entry.get("size", len(entry["data"]))
        dir_blob += struct.pack("<64sHBBIII16sIQQIQ", name, len(name), entry["type"], 1, entry["left"], entry["right"],
                                entry["child"], b"\0" * 16, 0, 0, 0, entry["start"], size)
    for _ in range(directory * 4 - len(entries)):
        dir_blob += struct.pack("<64sHBBIII16sIQQIQ", b"", 0, 0, 0, _FREE, _FREE, _FREE, b"\0" * 16, 0, 0, 0, 0, 0)
    body = pad(struct.pack(f"<{len(fat)}I", *fat)) + bytes(dir_blob)
    body += pad(struct.pack(f"<{len(mini_fat)}I", *mini_fat)) if mini_fat else b""
    body += pad(bytes(mini)) + b"".join(pad(entry["data"]) for entry in big)
    return header + body


def _properties(header: bytes, values: dict[tuple[int, int], object]) -> tuple[bytes, dict[str, bytes]]:
    streams: dict[str, bytes] = {}
    blob = bytearray(header)
    for (pid, ptyp), value in values.items():
        if ptyp == 0x001F:
            data = str(value).encode("utf-16-le")
            streams[f"__substg1.0_{pid:04X}{ptyp:04X}"] = data
            blob += struct.pack("<IIII", pid << 16 | ptyp, 6, len(data) + 2, 0)
        elif ptyp == 0x0102:
            data = bytes(value)  # type: ignore[arg-type]
            streams[f"__substg1.0_{pid:04X}{ptyp:04X}"] = data
            blob += struct.pack("<IIII", pid << 16 | ptyp, 6, len(data), 0)
        elif ptyp == 0x0040:
            blob += struct.pack("<IIQ", pid << 16 | ptyp, 6, int(value))  # type: ignore[call-overload]
        else:
            blob += struct.pack("<IIII", pid << 16 | ptyp, 6, int(value), 0)  # type: ignore[call-overload]
    streams["__properties_version1.0"] = bytes(blob)
    return bytes(blob), streams


def lzfu_literal(rtf: bytes) -> bytes:
    """RTF « compressé » sans référence (littéraux seuls) + marqueur de fin / literal-only LZFu + end marker."""
    body = bytearray()
    full = len(rtf) - len(rtf) % 8
    for start in range(0, full, 8):  # linéaire / linear
        body += b"\x00" + rtf[start:start + 8]
    rest = rtf[full:]
    # Bit de référence après les littéraux restants : marqueur de fin / reference bit after the rest: end marker.
    body += bytes([1 << len(rest)]) + rest + struct.pack(">H", ((207 + len(rtf)) & 0xFFF) << 4)
    return struct.pack("<IIII", len(body) + 12, len(rtf), 0x75465A4C, 0) + bytes(body)


def make_msg(*, body: str | None = None, html: bytes | None = None, rtf: bytes | None = None,
             message_class: str = "IPM.Note", codepage: int | None = None) -> bytes:
    filetime = (1_792_000_000 + 11_644_473_600) * 10_000_000
    root_values: dict[tuple[int, int], object] = {
        (0x001A, 0x001F): message_class, (0x0037, 0x001F): "Dossier NIR " + NIR,
        (0x0C1A, 0x001F): "Jean Dupont", (0x0C1F, 0x001F): EMAIL, (0x0039, 0x0040): filetime,
        (0x007D, 0x001F): f"From: Jean Dupont <{EMAIL}>\r\nTo: Marie Martin <marie.martin@exemple.fr>\r\n"
                          "Received: from mx (10.0.0.1)\r\nX-Originating-IP: [10.0.0.1]\r\n",
    }
    if body is not None:
        root_values[(0x1000, 0x001F)] = body
    if html is not None:
        root_values[(0x1013, 0x0102)] = html
    if rtf is not None:
        root_values[(0x1009, 0x0102)] = lzfu_literal(rtf)
    if codepage is not None:
        root_values[(0x3FDE, 0x0003)] = codepage  # PidTagInternetCodepage
    _, root = _properties(struct.pack("<8xIIII8x", 2, 1, 2, 1), root_values)
    _, to = _properties(b"\0" * 8, {(0x3001, 0x001F): "Marie Martin", (0x39FE, 0x001F): "marie.martin@exemple.fr",
                                    (0x0C15, 0x0003): 1})
    _, bcc = _properties(b"\0" * 8, {(0x3001, 0x001F): "Cache", (0x39FE, 0x001F): "cache@exemple.fr",
                                     (0x0C15, 0x0003): 3})
    _, attachment = _properties(b"\0" * 8, {(0x3707, 0x001F): "bulletin-dupont.pdf", (0x370E, 0x001F):
                                            "application/pdf", (0x3705, 0x0003): 1, (0x3701, 0x0102): b"%PDF-1.4"})
    tree = dict(root)
    tree["__recip_version1.0_#00000000"] = to
    tree["__recip_version1.0_#00000001"] = bcc
    tree["__attach_version1.0_#00000000"] = attachment
    tree["__nameid_version1.0"] = {"__substg1.0_00020102": b"", "__substg1.0_00030102": b"",
                                   "__substg1.0_00040102": b""}
    return make_cfb(tree)


# --- cibles du bac à sable / sandbox targets ---------------------------------------------------------


def echo(data: bytes, label: str) -> tuple[int, str, int]:
    return len(data), label, os.getpid()


def sleep_forever(data: bytes) -> None:
    while True:
        time.sleep(0.05)


def crash(data: bytes) -> None:
    os._exit(3)


def allocate(data: bytes) -> int:
    blocks = [bytearray(64 << 20) for _ in range(64)]  # 4 Gio : au-delà du plafond / above the cap
    return len(blocks)


def leak(data: bytes) -> None:
    raise ValueError("contenu secret " + data.decode())


def inspect_confinement(data: bytes) -> dict:
    """Ce qu'un convertisseur exploité verrait de son propre processus / what an exploited converter would see."""
    found: dict = {"secrets": sorted(k for k in os.environ if k.startswith("GDPRGUARDIAN_"))}
    if os.name == "posix":
        import ctypes
        import resource

        found["core"] = resource.getrlimit(resource.RLIMIT_CORE)[1]
        found["nproc"] = resource.getrlimit(resource.RLIMIT_NPROC)[1]
        found["stdout"] = os.readlink("/proc/self/fd/1") if os.path.exists("/proc/self/fd/1") else None
        if os.path.exists("/proc/self/status"):
            found["dumpable"] = ctypes.CDLL(None).prctl(3, 0, 0, 0, 0)  # PR_GET_DUMPABLE
    return found


def unserialisable(data: bytes) -> bytes:
    return data


COMPROMISED = "GDPRGUARDIAN_TEST_COMPROMISED"


def mark_compromised() -> None:
    os.environ[COMPROMISED] = "1"


class _Exploit:
    """Pickle qui exécute du code au chargement / pickle running code when loaded."""

    def __reduce__(self) -> tuple:
        return mark_compromised, ()


def forge_pickle_reply(data: bytes) -> None:
    """Travailleur compromis : envoie au parent un pickle piégé / compromised worker sends a trapped pickle."""
    import gc
    import pickle
    from multiprocessing.connection import _ConnectionBase  # type: ignore[attr-defined]

    for candidate in gc.get_objects():
        if isinstance(candidate, _ConnectionBase) and not candidate.closed:
            candidate.send_bytes(pickle.dumps(("ok", _Exploit())))
            return
