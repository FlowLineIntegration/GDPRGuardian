"""Lecture xlsx en Rust contre la référence python-calamine ; pipeline tableur natif contre pur.

Rust xlsx reading against the python-calamine reference; native against pure spreadsheet pipeline.

Trois niveaux de preuve / three levels of proof :
- feuilles : classeurs XML générés (Hypothesis) couvrant chaque type de cellule, style, date 1900/1904, chaîne
  partagée, enrichie, en ligne, phonétique, formule, fusion, préfixe d'espace de noms, entités et échappements
  `_xHHHH_`, plus des formes invalides. Quand le noyau accepte, valeurs, types, étendue et fusions sont
  identiques à calamine ; quand calamine échoue, le noyau refuse toujours. Sheets: generated XML workbooks;
  when the core accepts, values, types, extent and merges equal calamine's; when calamine fails, the core
  always declines.
- garde : chemin rapide de `xlsx_extent` et comptes de formules et d'en-têtes identiques en Rust, sur une partie
  entière comme pendant la lecture en flux ; état de la garde identique / guard twins, whole part or streamed.
- bout en bout : corpus de classeurs et CSV, fichiers hostiles ; sortie, rapport, avertissements et éléments
  écartés identiques avec et sans noyau / end to end: identical output, report, warnings and dropped parts.
"""

from __future__ import annotations

import datetime as dt
import io
import random
import sys
import zipfile
from pathlib import Path
from typing import Any

import pytest
from hypothesis import HealthCheck, given, settings
from hypothesis import strategies as st

from gdprguardian import Anonymizer, _native
from gdprguardian.documents import DocumentOptions
from gdprguardian.tabular import _package, _reader, anonymize_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent / "tabular"))

from tabular_corpus import csv_bytes, workbook_bytes
from tabular_factories import FAR_SHEET, raw_xlsx, rezip, xlsx_bytes, zip_bytes

_CORE = _native.core()
pytestmark = pytest.mark.skipif(_CORE is None or not hasattr(_CORE, "XlsxBook"),
                                reason="noyau Rust absent / no Rust core")
_FUZZ = settings(max_examples=300, deadline=None, suppress_health_check=[HealthCheck.too_slow])

MAIN = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RELS = "http://schemas.openxmlformats.org/package/2006/relationships"
DOC_RELS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


# --- vues comparables des deux lecteurs / comparable views of both readers ---------------------------------------


def _calamine(data: bytes) -> Any:
    from python_calamine import CalamineWorkbook

    try:
        book = CalamineWorkbook.from_filelike(io.BytesIO(data))
    except BaseException as exc:  # toute erreur, panique comprise / any error, panic included
        return ("open-error", type(exc).__name__)
    metadata, sheets = [], {}
    for meta in book.sheets_metadata:
        kind = "worksheet" if str(meta.typ).endswith("WorkSheet") else str(meta.typ).rsplit(".", 1)[1].lower()
        metadata.append((meta.name, kind, str(meta.visible).rsplit(".", 1)[1].lower()))
        try:
            sheet = book.get_sheet_by_name(meta.name)
            rows = sheet.to_python()
            sheets.setdefault(meta.name, (sheet.height, sheet.width, sheet.start, sheet.merged_cell_ranges or [],
                                          _typed(rows)))
        except BaseException as exc:
            sheets.setdefault(meta.name, ("error", type(exc).__name__))
    return metadata, sheets


def _typed(rows: list[list[Any]]) -> list[list[tuple[str, Any]]]:
    """Valeurs avec leur type exact ; NaN comparé par son texte / values with their exact type."""
    return [[(type(value).__name__, repr(value) if value != value else value) for value in row] for row in rows]


def _compare(data: bytes) -> str:
    """« declined », « partial » ou « same » ; échoue sur toute divergence / fails on any divergence."""
    reference = _calamine(data)
    try:
        book = _CORE.XlsxBook(data)
    except _CORE.XlsxDecline:
        return "declined"
    assert reference[0] != "open-error", f"calamine refuse, le noyau accepte : {reference}"
    metadata, sheets = reference
    assert book.sheets == metadata
    outcome = "same"
    for name, kind, _ in book.sheets:
        if kind != "worksheet":
            continue
        try:
            sheet = book.sheet(name)
            native = (sheet.height, sheet.width, sheet.start, sheet.merged, _typed(sheet.rows()))
        except _CORE.XlsxDecline:
            outcome = "partial"
            continue
        except OverflowError:  # même erreur que calamine (durée hors bornes) / same error as calamine
            assert sheets[name] == ("error", "OverflowError")
            continue
        assert sheets[name] == native, name
    return outcome


# --- classeurs générés / generated workbooks ---------------------------------------------------------------------

_CANONICAL_NUMBERS = st.sampled_from([
    "0", "1", "-1", "1.5", "1e5", "1E-05", "-0", "007", "45000", "45000.25", "0.75", "59", "60", "61", "366", "1e12",
    "-1e12", "2958465.99999", "1e400", "-2e11", "12345678901234567890", "0.1", "1462", "1461.5", "-0.25"]) | (
    st.floats(allow_nan=False, allow_infinity=False).map(repr).filter(lambda s: "inf" not in s))
_NUMBERS = _CANONICAL_NUMBERS | st.sampled_from(["1.", ".5", "+1", " 1", "nan", "inf", "-inf", "abc", "1,5", "0x10"])
_TEXT_PIECES = ["a", "Zé", " ", " ", "\t", "\r\n", "\n", "\r", "_x000D_", "_x005F_x000D_", "_x00+5_", "&amp;",
                "&lt;", "&#233;", "&#x1F600;", "😀", "é", "\u001c", "&#13;", "x" * 40, "[PERSONNE_1]", "|"]
_BAD_PIECES = ["&", "&nbsp;", "<!-- c -->", "<![CDATA[x]]>", "<?pi x?>", "&#0;", "<b>", "&#xD800;"]
_TEXT = st.lists(st.sampled_from(_TEXT_PIECES), max_size=6).map("".join)
_RISKY_TEXT = st.lists(st.sampled_from(_TEXT_PIECES + _BAD_PIECES), max_size=4).map("".join)
_FORMATS = ["dd/mm/yyyy", "[h]:mm:ss", "hh:mm", "yyyy-mm-dd hh:mm:ss", "0.00", "#,##0 €", "\"Y\" 0", "[ss].0",
            "mmm-yy;@", "General", "d\"jour\"", "[Red]0.0;[Blue]-0.0", "a/p h"]
_BUILTINS = ["0", "1", "14", "15", "18", "20", "22", "45", "46", "47", "49", "014"]


def _ref(row: int, col: int) -> str:
    letters = ""
    col += 1
    while col:
        col, rest = divmod(col - 1, 26)
        letters = chr(65 + rest) + letters
    return f"{letters}{row + 1}"


@st.composite
def _cell(draw: Any, row: int, col: int, strings: int, styles: int, prefix: str, canonical: bool) -> str:
    """Cellule canonique (forme d'Excel, LibreOffice, openpyxl…) ou retorse / canonical or tricky cell."""
    p = prefix
    kinds = [None, "n", "s", "b", "e", "str", "inlineStr", None, "n"] if canonical else [None, "s", "d", "zz", "e"]
    kind = draw(st.sampled_from(kinds))
    if kind == "s" and strings == 0 and canonical:
        kind = "n"
    attrs = []
    position = draw(st.sampled_from(["ref", "ref", "ref", "none", "lower"] + ([] if canonical else ["bad"] * 3)))
    if position == "ref":
        attrs.append(("r", _ref(row, col)))
    elif position == "lower":
        attrs.append(("r", _ref(row, col).lower()))
    elif position == "bad":
        attrs.append(("r", draw(st.sampled_from(["A0", "1A", "A", "A1:B2", "XFE1", ""]))))
    if draw(st.booleans()):
        styles_seen = [str(i) for i in range(styles + 2)]
        attrs.append(("s", draw(st.sampled_from(styles_seen if canonical else [*styles_seen, "01", "", "x"]))))
    if kind is not None:
        attrs.append(("t", kind))
    if not canonical and draw(st.integers(0, 4)) == 0 and attrs:
        attrs.append(attrs[0])  # attribut en double / duplicate attribute
    order = draw(st.permutations(attrs))
    quote = draw(st.sampled_from(['"', "'"]))
    spaces = draw(st.sampled_from([" ", " ", "\t", "\n  "]))
    attr_text = "".join(f"{spaces}{key}={quote}{value}{quote}" for key, value in order)
    if kind == "s":
        valid = [str(i) for i in range(strings)]
        value = draw(st.sampled_from(valid if canonical else [*valid, str(strings), "01", "-1", " 1"]))
    elif kind == "b":
        value = draw(st.sampled_from(["0", "1"] if canonical else ["true", "2", "0"]))
    elif kind == "e":
        errors = ["#DIV/0!", "#N/A", "#NAME?", "#NULL!", "#NUM!", "#REF!", "#VALUE!"]
        value = draw(st.sampled_from(errors if canonical else [*errors, "#SPILL!", "x"]))
    elif kind in ("str", "zz"):
        value = draw(_TEXT if canonical else _RISKY_TEXT)
    elif kind == "d":
        value = draw(st.sampled_from(["2024-01-02", "2024-01-02T03:04:05", "03:04:05", "x"]))
    else:
        value = draw(_CANONICAL_NUMBERS if canonical else _NUMBERS)
    inline = f"<{p}is><{p}t>{draw(_TEXT)}</{p}t></{p}is>"
    rich = f"<{p}is><{p}r><{p}t>{draw(_TEXT)}</{p}t></{p}r><{p}r><{p}rPr><{p}b/></{p}rPr><{p}t>z</{p}t></{p}r></{p}is>"
    forms = [[f"<{p}v>{value}</{p}v>"], [f"<{p}v>{value}</{p}v>"], [f"<{p}f>A1+1</{p}f>", f"<{p}v>{value}</{p}v>"],
             [f"<{p}v>{value}</{p}v>", f"<{p}f>SUM(A1:A2)</{p}f>"], [inline], [rich], [f"<{p}v/>"],
             [f"<{p}v></{p}v>"], [], [f"<{p}f t=\"shared\" si=\"0\"/><{p}v>{value}</{p}v>"]]
    if not canonical:
        forms += [[f"<{p}v>{value}</{p}v>", inline], [f"<{p}x>1</{p}x>"], [f"<{p}v>{draw(_RISKY_TEXT)}</{p}v>"]]
    children = draw(st.sampled_from(forms))
    if not children and draw(st.booleans()):
        return f"<{p}c{attr_text}/>"
    return f"<{p}c{attr_text}>{''.join(children)}</{p}c>"


@st.composite
def _sheet(draw: Any, strings: int, styles: int, prefix: str, canonical: bool) -> str:
    p = prefix
    rows = []
    for row in range(draw(st.integers(0, 8))):
        row_number = row + draw(st.integers(0, 2))
        cells = [draw(_cell(row_number, col + draw(st.integers(0, 1)), strings, styles, p, canonical))
                 for col in range(draw(st.integers(0, 6)))]
        r = draw(st.sampled_from([f' r="{row_number + 1}"', "", f' r="{row_number + 1}" spans="1:5" ht="15"',
                                  f' customHeight="1" r="{row_number + 1}"']))
        rows.append(f"<{p}row{r}/>" if not cells and draw(st.booleans()) else f"<{p}row{r}>{''.join(cells)}</{p}row>")
    heads = ["", f'<{p}dimension ref="A1:C3"/>', f'<{p}dimension ref="A1"/>',
             f"<{p}sheetViews><{p}sheetView workbookViewId=\"0\"/></{p}sheetViews>"]
    head = draw(st.sampled_from(heads if canonical else [*heads, f'<{p}dimension ref="bad"/>']))
    refs = ["A1:B2", "C3:C5", "B1:A1", "A1", "Z9:AA10"]
    merges = draw(st.lists(st.sampled_from(refs if canonical else [*refs, "bad", "A1:B2:C3"]), max_size=3))
    tail = ""
    if merges:
        tail = f"<{p}mergeCells count=\"{len(merges)}\">" + "".join(
            f'<{p}mergeCell ref="{ref}"/>' for ref in merges) + f"</{p}mergeCells>"
    declaration = draw(st.sampled_from(['<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n', "",
                                        "<?xml version='1.0' encoding='utf-8'?>", "﻿<?xml version=\"1.0\"?>"]))
    namespace = f'xmlns:{p[:-1]}="{MAIN}"' if p else f'xmlns="{MAIN}"'
    data = f"<{p}sheetData>{''.join(rows)}</{p}sheetData>" if canonical or draw(st.integers(0, 5)) else ""
    return f"{declaration}<{p}worksheet {namespace}>{head}{data}{tail}</{p}worksheet>"


@st.composite
def _workbooks(draw: Any) -> bytes:
    canonical = draw(st.integers(0, 9)) < 7
    prefix = draw(st.sampled_from(["", "", "x:"]))
    forms = [f"<si><t>{draw(_TEXT)}</t></si>", f'<si><t xml:space="preserve">{draw(_TEXT)}</t></si>',
             f"<si><r><rPr><b/></rPr><t>{draw(_TEXT)}</t></r><r><t xml:space=\"preserve\"> {draw(_TEXT)}</t></r></si>",
             "<si><t>base</t><rPh sb=\"0\" eb=\"1\"><t>phon</t></rPh><phoneticPr fontId=\"1\"/></si>", "<si/>",
             "<si><t/></si>"]
    if not canonical:
        forms.append(f"<si><t>{draw(_RISKY_TEXT)}</t></si>")
    shared = [draw(st.sampled_from(forms)) for _ in range(draw(st.integers(0, 4)))]
    formats = draw(st.lists(st.sampled_from(_FORMATS), max_size=4))
    xfs = draw(st.lists(st.sampled_from(_BUILTINS + [str(164 + i) for i in range(len(formats))] + [None]),
                        max_size=6))
    styles = ('<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="' + MAIN + '"><numFmts count="'
              + str(len(formats)) + '">' + "".join(
                  f'<numFmt numFmtId="{164 + i}" formatCode="{code.replace(chr(34), "&quot;")}"/>'
                  for i, code in enumerate(formats)) + "</numFmts><cellXfs>" + "".join(
                  "<xf/>" if xf is None else f'<xf numFmtId="{xf}" fontId="0"/>' for xf in xfs)
              + "</cellXfs></styleSheet>")
    names = draw(st.lists(st.sampled_from(["Feuil1", "Données &amp; co", "A&#9;B", "Été", "Feuil1"]), min_size=1,
                          max_size=3))
    states = [draw(st.sampled_from(["", ' state="hidden"', ' state="veryHidden"', ' state="visible"']))
              for _ in names]
    date1904 = draw(st.sampled_from(["", '<workbookPr date1904="1"/>', '<workbookPr date1904="true"/>',
                                     '<workbookPr date1904="TRUE"/>', '<workbookPr date1904="0"/>']))
    sheets = "".join(f'<sheet name="{name}" sheetId="{i + 1}" r:id="rId{i + 1}"{state}/>'
                     for i, (name, state) in enumerate(zip(names, states, strict=True)))
    workbook = (f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="{MAIN}" xmlns:r="{DOC_RELS}">{date1904}'
                f'<sheets>{sheets}</sheets><definedNames><definedName name="zone">Feuil1!$A$1</definedName>'
                "</definedNames></workbook>")
    rels = "".join(f'<Relationship Id="rId{i + 1}" Type="{DOC_RELS}/worksheet" Target="worksheets/sheet{i + 1}.xml"/>'
                   for i in range(len(names)))
    members = [
        ("[Content_Types].xml", b"<Types/>"),
        ("_rels/.rels", (f'<Relationships xmlns="{RELS}"><Relationship Id="rId1" Type="{DOC_RELS}/officeDocument" '
                         'Target="xl/workbook.xml"/></Relationships>').encode()),
        ("xl/workbook.xml", workbook.encode()),
        ("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{RELS}">{rels}</Relationships>'.encode()),
        ("xl/styles.xml", styles.encode()),
    ]
    if shared or draw(st.booleans()):
        members.append(("xl/sharedStrings.xml",
                        f'<sst xmlns="{MAIN}" uniqueCount="{len(shared)}">{"".join(shared)}</sst>'.encode()))
    members.extend((f"xl/worksheets/sheet{i + 1}.xml", draw(_sheet(len(shared), len(xfs), prefix, canonical)).encode())
                   for i in range(len(names)))
    return zip_bytes(members)


@_FUZZ
@given(_workbooks())
def test_generated_workbooks_read_exactly_as_calamine(data: bytes) -> None:
    _compare(data)


def test_canonical_workbooks_are_never_declined() -> None:
    """Le corpus (xlsxwriter : textes, nombres, dates, durées, booléens, formules, enrichis, fusions, tableaux)
    est lu en natif sans refus, à l'identique / the corpus is read natively, never declined, identically."""
    outcomes = [_compare(workbook_bytes(seed)) for seed in range(80)]
    assert outcomes == ["same"] * len(outcomes)


def test_openpyxl_workbooks_inline_strings_and_1904() -> None:
    import openpyxl

    for epoch in (openpyxl.utils.datetime.CALENDAR_WINDOWS_1900, openpyxl.utils.datetime.CALENDAR_MAC_1904):
        workbook = openpyxl.Workbook()
        workbook.epoch = epoch
        sheet = workbook.active
        sheet.append(["Nom", "Montant", "Date", "Vrai", None, "Heure", "Durée"])
        for i in range(1, 30):
            sheet.append([f"Nom {i} é", i / 7, dt.datetime(2020, 1, 1) + dt.timedelta(days=i, hours=i), i % 3 == 0,
                          None, dt.time(i % 24, i, i), dt.timedelta(hours=i * 7, minutes=i)])
        sheet["J5"] = "=1/0"
        sheet["K9"] = dt.date(1904, 1, 1)
        sheet.merge_cells("L1:M4")
        buffer = io.BytesIO()
        workbook.save(buffer)
        assert _compare(buffer.getvalue()) == "same"


def test_inline_strings_written_by_hand_match() -> None:
    cells = ('<row r="1"><c r="A1" t="inlineStr"><is><t xml:space="preserve"> a b </t></is></c>'
             '<c r="B1" t="inlineStr"><is><r><t>x</t></r><r><t>y</t></r></is></c>'
             '<c r="C1" t="str"><f>A1</f><v>  calc &amp; co\r\n</v></c><c r="D1" t="b"><v>0</v></c>'
             '<c r="E1" t="e"><v>#N/A</v></c><c r="F1"><v>1e-3</v></c></row>')
    sheet = (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{MAIN}"><sheetData>{cells}</sheetData>'
             "</worksheet>").encode()
    data = raw_xlsx('<sheet name="S" sheetId="1" r:id="rId1"/>', {"rId1": "worksheets/sheet1.xml"},
                    [("xl/worksheets/sheet1.xml", sheet)])
    assert _compare(data) == "same"


@pytest.mark.parametrize("bad", [
    b"<!-- commentaire --><worksheet/>", b"<worksheet><sheetData><row><c><v>1</v></c></row></sheetData>",
    "<?xml version=\"1.0\" encoding=\"ISO-8859-1\"?><worksheet><sheetData/></worksheet>".encode("latin-1"),
    b"\xff\xfe<\x00w\x00", b"<worksheet><sheetData><c r=\"A1\"><v>&nbsp;</v></c></sheetData></worksheet>",
])
def test_non_canonical_sheets_are_declined_or_identical(bad: bytes) -> None:
    data = raw_xlsx('<sheet name="S" sheetId="1" r:id="rId1"/>', {"rId1": "worksheets/sheet1.xml"},
                    [("xl/worksheets/sheet1.xml", bad)])
    assert _compare(data) in ("declined", "partial", "same")


# --- garde : jumeaux du chemin rapide / guard: fast-path twins ---------------------------------------------------

_EXTENT_PIECES = ['<row r="1">', '<row r="250000">', "</row>", '<c r="A1">', '<c r="ZZ9">', '<c r="XFD1048576">',
                  '<c r="AB123456">', "<c>", "<c/>", '<c\tr="B2" s="1"/>', "<c r='C3'>", '<c r="a1">',
                  '<c r="1">', '<c r="A1x">', '<c r="AAAAAAAAAAAAAAAAA1">', "<v>1</v>", "</c>", " r ", " r= ",
                  "<!-- -->", "<?pi?>", "<x:c>", ":row", "<f>", "<f t=\"shared\">", ":f>", "Header>", " ",
                  '<row r = "5">', '<c r="A12345678901234567">', "\xef\xbb\xbf", "\n", "'", '"']
_EXTENT_XML = st.lists(st.sampled_from(_EXTENT_PIECES), max_size=20).map(lambda p: "".join(p).encode("latin-1"))


@_FUZZ
@given(_EXTENT_XML, st.sampled_from([1 << 24, 100, 10**9]))
def test_extent_twin_matches_python(xml: bytes, limit: int) -> None:
    assert _package.xlsx_extent(xml, limit, native=True) == _package.xlsx_extent(xml, limit)
    assert _package.sheet_measure(xml, limit, native=True) == _package.sheet_measure(xml, limit)
    formulas = xml.count(b"<f>") + xml.count(b"<f ") + xml.count(b":f>") + xml.count(b":f ")
    assert _CORE.sheet_measure(xml, limit)[1:] == (formulas, b"Header>" in xml or b"Footer>" in xml)


def test_extent_twin_on_real_and_utf16_sheets() -> None:
    for seed in range(20):
        archive = zipfile.ZipFile(io.BytesIO(workbook_bytes(seed)))
        for name in archive.namelist():
            if name.startswith("xl/worksheets/"):
                xml = archive.read(name)
                assert _package.sheet_measure(xml, native=True) == _package.sheet_measure(xml)
    utf16 = b"\xff\xfe" + FAR_SHEET.decode().replace('encoding="UTF-8"', 'encoding="UTF-16"').encode("utf-16-le")
    assert _CORE.sheet_measure(utf16, 1 << 24)[0] is None
    assert _package.xlsx_extent(utf16, native=True) == _package.xlsx_extent(utf16)


def _streamed_measures(data: bytes, limit: int) -> list[Any]:
    """Mesures prises par le noyau pendant la lecture, et celles de la garde sur les mêmes parties.
    Measurements taken by the core while reading, and the guard's on the same parts."""
    try:
        book = _CORE.XlsxBook(data)
    except _CORE.XlsxDecline:
        return []
    archive = zipfile.ZipFile(io.BytesIO(data))
    infos = {info.filename.lower(): info for info in archive.infolist()}
    pairs = []
    for name, kind, _ in book.sheets:
        try:
            measured = book.sheet(name).measured(limit) if kind == "worksheet" else None
        except _CORE.XlsxDecline:
            continue
        if measured is None:
            continue
        part, size, crc, extent, formulas, headers = measured
        info = infos[part]
        xml = archive.read(info)
        expected = (info.filename.lower(), info.file_size, info.CRC, *_package.sheet_measure(xml, limit))
        pairs.append(((part, size, crc, extent or expected[3], formulas, headers), expected))
    return pairs


_BALANCED_ROWS = ['<row r="1"/>', '<row r="250000"><c r="B2"><v>1</v></c></row>', "<row><c><v>2</v></c></row>",
                  '<row r="3"><c r="ZZ3"><v>0</v></c><c r="AB123456"/></row>',
                  '<row r="4"><c r="C4"><f>A1</f><v>5</v></c><c r="D4" t="str"><f t="shared">B1</f><v>x</v></c></row>',
                  '<row r="5" spans="1:2"><c r="XFD5"><v>1</v></c></row>', "<row\tr='6'><c\nr='D6'><v>1</v></c></row>",
                  '<row r="7"><c r="E7" t="inlineStr"><is><t>a r b Header&gt;</t></is></c></row>',
                  '<row r="8"><c r="AAA8" s="0"><v>3</v></c></row>', '<row r="9"><c r="A9" t="n"><v>1.5</v></c></row>']
_TAILS = ["", '<mergeCells count="1"><mergeCell ref="A1:B2"/></mergeCells>',
          "<headerFooter><oddHeader>&amp;L x</oddHeader><oddFooter>y</oddFooter></headerFooter>"]


@_FUZZ
@given(st.lists(st.sampled_from(_BALANCED_ROWS), max_size=12), st.sampled_from(_TAILS),
       st.sampled_from([1 << 24, 100]))
def test_measure_taken_while_reading_matches_the_guard(rows: list[str], tail: str, limit: int) -> None:
    sheet = (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{MAIN}"><sheetData>{"".join(rows)}'
             f"</sheetData>{tail}</worksheet>").encode()
    data = raw_xlsx('<sheet name="S" sheetId="1" r:id="rId1"/>', {"rId1": "worksheets/sheet1.xml"},
                    [("xl/worksheets/sheet1.xml", sheet)])
    for streamed, expected in _streamed_measures(data, limit):
        assert streamed == expected


def test_measure_taken_while_reading_matches_on_the_corpus() -> None:
    compared = 0
    # Une graine sur dix : feuilles de plus de 256 Kio, la mesure traverse les morceaux du lecteur.
    # One seed in ten: sheets over 256 KiB, the measurement spans the reader's chunks.
    for seed in range(30):
        data = workbook_bytes(seed, rows=150 if seed % 10 else 8000)
        for streamed, expected in _streamed_measures(data, 1 << 24):
            assert streamed == expected
            compared += 1
    assert compared >= 30


def _guard_state(data: bytes, native: bool) -> Any:
    package = _package.open_package(data, "classeur.xlsx", native=native)
    try:
        read = _reader.read_workbook(data, package, DocumentOptions(include_hidden=True), lambda: None, native)
    finally:
        package.close()
    return (package.measured, dict(package.dropped), package.warnings, read.skipped,
            [_typed(sheet.rows) for sheet in read.sheets])


def test_guard_state_is_identical_when_the_core_measures() -> None:
    hostile = _hostile()
    cases = [*(workbook_bytes(seed) for seed in range(12)), hostile["far"], hostile["spill-error"],
             hostile["overlapping-entries"]]
    for data in cases:
        try:
            expected = _guard_state(data, native=False)
        except Exception as exc:
            with pytest.raises(type(exc)):
                _guard_state(data, native=True)
            continue
        assert _guard_state(data, native=True) == expected


def test_guard_reuses_the_measurement_taken_while_reading(monkeypatch: pytest.MonkeyPatch) -> None:
    """Avec le noyau, la garde ne décompresse aucune feuille une seconde fois / no sheet is inflated twice."""
    read = _package.Package.read
    sheets_read: list[str] = []

    def spy(self: _package.Package, name: str, limit: int = _package.MAX_PART) -> bytes | None:
        if name.startswith("xl/worksheets/") and "/_rels/" not in name:
            sheets_read.append(name)
        return read(self, name, limit)

    monkeypatch.setattr(_package.Package, "read", spy)
    for seed in (0, 10):
        _guard_state(workbook_bytes(seed, rows=400), native=True)
    assert sheets_read == []
    _guard_state(workbook_bytes(0, rows=400), native=False)
    assert sheets_read  # la référence mesure elle-même / the reference measures by itself


# --- bout en bout : natif contre pur / end to end: native against pure --------------------------------------------

_NATIVE = Anonymizer(countries=("FR",))
_PURE = Anonymizer(countries=("FR",), native=False)


def _outcome(data: bytes, filename: str, engine: Anonymizer, **options: Any) -> Any:
    try:
        result = anonymize_workbook(data, filename, engine, DocumentOptions(**options))
    except Exception as exc:
        return "error", type(exc).__name__, str(exc)
    report = {key: value for key, value in result.report.items() if key != "timings_ms"}
    return result.markdown, result.kind, report, result.warnings, result.dropped


@pytest.mark.parametrize("seed", range(24))
def test_corpus_output_is_identical_with_and_without_the_core(seed: int) -> None:
    variants = [{"sandbox": False}, {"sandbox": False, "include_hidden": True, "table_format": "csv"},
                {"sandbox": False, "max_cells": 70}, {"sandbox": False, "max_rows_per_sheet": 3}]
    for options in variants[:1] if seed % 3 else variants:
        for data, filename in ((workbook_bytes(seed), "classeur.xlsx"), (csv_bytes(seed), "export.csv")):
            assert _outcome(data, filename, _NATIVE, **options) == _outcome(data, filename, _PURE, **options)


def test_sandboxed_reading_agrees_with_the_core() -> None:
    data = workbook_bytes(3)
    expected = _outcome(data, "classeur.xlsx", _PURE, sandbox=False)
    assert _outcome(data, "classeur.xlsx", _NATIVE, sandbox=True) == expected
    assert _outcome(data, "classeur.xlsx", _NATIVE, sandbox=False) == expected


def _bomb() -> bytes:
    return zip_bytes([("[Content_Types].xml", b"<Types/>"), ("xl/workbook.xml", b"<w/>"),
                      ("xl/worksheets/sheet1.xml", b"\0" * (4 << 20))])


def _overlapping(data: bytes, name: bytes) -> bytes:
    """Taille compressée annoncée qui déborde sur l'entrée suivante (`zipfile` : « Overlapped entries »).
    Declared compressed size running into the next entry."""
    at = data.index(name, data.index(b"PK\x01\x02")) - 46
    size = int.from_bytes(data[at + 20:at + 24], "little")
    return data[:at + 20] + (size + 4).to_bytes(4, "little") + data[at + 24:]


def _hostile() -> dict[str, bytes]:
    simple = xlsx_bytes(lambda wb: wb.add_worksheet("S").write("A1", "Dupont"))

    def far(wb: Any) -> None:
        sheet = wb.add_worksheet("Loin")
        sheet.write(0, 0, "a")
        sheet.write(1_048_575, 16_383, "z")
        wb.add_worksheet("Près").write_row(0, 0, ["Email", "x@exemple.fr"])

    return {
        "bomb": _bomb(),
        "far": xlsx_bytes(far),
        "broken-xml": rezip(simple, replace={"xl/worksheets/sheet1.xml": b"<worksheet><sheetData><row>"}),
        "broken-workbook": rezip(simple, replace={"xl/workbook.xml": b"<workbook><sheets><sheet"}),
        "no-rels": rezip(simple, replace={"xl/_rels/workbook.xml.rels": b"<Relationships/>"}),
        "wrong-types": rezip(simple, replace={"[Content_Types].xml": b"<Types><Default Extension=\"xml\" "
                                                                      b"ContentType=\"text/plain\"/></Types>"}),
        "dtd": rezip(simple, replace={"xl/workbook.xml": lambda raw: raw.replace(
            b"?>", b'?><!DOCTYPE w [<!ENTITY a "aaaa">]>', 1)}),
        "truncated": simple[: len(simple) // 2],
        "bad-shared-index": rezip(simple, replace={"xl/worksheets/sheet1.xml": lambda raw: raw.replace(
            b"<v>0</v>", b"<v>99</v>")}),
        "spill-error": rezip(simple, replace={"xl/worksheets/sheet1.xml": lambda raw: raw.replace(
            b't="s"><v>0</v>', b't="e"><v>#SPILL!</v>')}),
        "huge-duration": xlsx_bytes(lambda wb: wb.add_worksheet("D").write_number(
            "A1", 1e12, wb.add_format({"num_format": "[h]:mm:ss"}))),
        "negative-date": xlsx_bytes(lambda wb: wb.add_worksheet("D").write_number(
            "A1", -3e11, wb.add_format({"num_format": "dd/mm/yyyy"}))),
        "overlapping-entries": _overlapping(simple, b"xl/worksheets/sheet1.xml"),
    }


@pytest.mark.parametrize("name", sorted(_hostile()))
def test_hostile_files_are_handled_identically(name: str) -> None:
    data = _hostile()[name]
    assert _outcome(data, "x.xlsx", _NATIVE, sandbox=False) == _outcome(data, "x.xlsx", _PURE, sandbox=False)


def test_python_reference_is_kept_without_the_core() -> None:
    data = workbook_bytes(1)
    package = _package.open_package(data, "x.xlsx")
    try:
        assert type(_reader._open_book(data, package, native=False)).__name__ == "_CalamineBook"
        assert type(_reader._open_book(data, package, native=True)).__name__ == "_NativeBook"
    finally:
        package.close()


# --- vérification indépendante : autres générateurs / independent check: other generators -----------------------


def _one_cell_book(value: str, num_fmt: str, date1904: bool) -> bytes:
    """Une cellule numérique de style `num_fmt` (identifiant intégré ou code personnalisé) / one styled number."""
    builtin = num_fmt.isdigit()
    styles = (f'<?xml version="1.0" encoding="UTF-8"?><styleSheet xmlns="{MAIN}">'
              + ("" if builtin else f'<numFmts count="1"><numFmt numFmtId="164" formatCode="{num_fmt}"/></numFmts>')
              + f'<cellXfs count="1"><xf numFmtId="{num_fmt if builtin else 164}"/></cellXfs></styleSheet>')
    workbook = (f'<?xml version="1.0" encoding="UTF-8"?><workbook xmlns="{MAIN}" xmlns:r="{DOC_RELS}">'
                + ('<workbookPr date1904="1"/>' if date1904 else "")
                + '<sheets><sheet name="S" sheetId="1" r:id="rId1"/></sheets></workbook>')
    sheet = (f'<?xml version="1.0" encoding="UTF-8"?><worksheet xmlns="{MAIN}"><sheetData><row r="1">'
             f'<c r="B2" s="0"><v>{value}</v></c></row></sheetData></worksheet>')
    rels = f'<Relationship Id="rId1" Type="{DOC_RELS}/worksheet" Target="worksheets/sheet1.xml"/>'
    return zip_bytes([
        ("[Content_Types].xml", b"<Types/>"),
        ("_rels/.rels", (f'<Relationships xmlns="{RELS}"><Relationship Id="rId1" Type="{DOC_RELS}/officeDocument" '
                         'Target="xl/workbook.xml"/></Relationships>').encode()),
        ("xl/workbook.xml", workbook.encode()),
        ("xl/_rels/workbook.xml.rels", f'<Relationships xmlns="{RELS}">{rels}</Relationships>'.encode()),
        ("xl/styles.xml", styles.encode()),
        ("xl/worksheets/sheet1.xml", sheet.encode()),
    ])


def _serial(year: int, month: int, day: int) -> int:
    """Numéro de série Excel d'une date du calendrier grégorien proleptique, toute année (H. Hinnant).
    Excel serial of a proleptic Gregorian date, any year."""
    def days(y: int, m: int, d: int) -> int:
        y -= m <= 2
        era = y // 400
        yoe = y - era * 400
        doy = (153 * (m + (-3 if m > 2 else 9)) + 2) // 5 + d - 1
        return era * 146_097 + yoe * 365 + yoe // 4 - yoe // 100 + doy

    return days(year, month, day) - days(1899, 12, 30)


def test_date_and_duration_bounds_read_exactly_as_calamine() -> None:
    """Bornes de chrono (années ±262 143), de Python (1 à 9 999), des millisecondes i64 et de `timedelta`, bogue
    de 1900, système 1904 : valeur, type et erreur identiques à calamine, sinon refus. Chrono, Python, i64
    milliseconds and timedelta bounds, 1900 bug, 1904 system: same value, type and error as calamine, or decline."""
    values = [repr(float(_serial(*date) + delta))
              for date in [(-262_144, 12, 31), (-262_143, 1, 1), (262_142, 12, 31), (262_143, 1, 1), (1, 1, 1),
                           (9_999, 12, 31), (10_000, 1, 1), (1900, 2, 28), (1900, 3, 1)]
              for delta in (-1, -0.5, 0, 0.25, 0.9999999, 1)]
    ms_days = (2**63 - 1) / 86_400_000
    values += [repr(v) for v in (ms_days, -ms_days, ms_days * 1.000001, -ms_days * 1.000001, 999_999_999.5,
                                 -999_999_999.5, 1e9 + 1, 1e15, -1e15, 0.0, -0.0, 0.5, 59.0, 60.0, 61.0, -1462.0)]
    values += ["1e400", "-1e400", "1e-400"]
    for value in values:
        for num_fmt in ("14", "22", "46", "[h]:mm:ss", "hh:mm"):
            for date1904 in (False, True):
                assert _compare(_one_cell_book(value, num_fmt, date1904)) in ("same", "partial"), (value, num_fmt)


def test_random_number_formats_are_classified_as_calamine() -> None:
    """Codes de format tirés au hasard (guillemets, échappements, crochets, AM/PM, sections) : date, durée ou
    nombre, comme calamine / random format codes: date, duration or number, as calamine."""
    rng = random.Random(1789)
    alphabet = [*'dmyhsDMYHSAPap/[]_\\*";:.0#@ E+-%é', "&quot;", "&amp;", "AM/PM", "[h]", "[mm]", "[ss]", "[Red]"]
    for _ in range(250):
        code = "".join(rng.choice(alphabet) for _ in range(rng.randint(0, 10))).replace('"', "&quot;")
        for value in ("45000.5", "0.25", "-3.75"):
            assert _compare(_one_cell_book(value, code, False)) == "same", code


_MUTATION_TOKENS = [
    b'<c r="B2"><v>1</v></c>', b'<c r="C3" t="s"><v>0</v></c>', b"&amp;", b"&#233;", b"&#x1F600;", b"&#0;",
    b"&nbsp;", b"<v>", b"</v>", b'<row r="7">', b"</row>", b"<c>", b"</c>", b' t="str"', b' t="e"', b' t="d"',
    b' s="2"', b' r="ZZ99"', b"<is><t>x</t></is>", b"<f>A1</f>", b'<mergeCell ref="A1:B3"/>', b"<!-- c -->",
    b"<![CDATA[x]]>", b"_x000D_", b"\r\n", b"\r", b"1e400", b"-1e400", b"1.", b"+1", b' xml:space="preserve"',
    b"<rPh><t>p</t></rPh>", b"<v/>", b"\xef\xbb\xbf", b"\xf0\x9f\x98\x80", b"\xcc\x81", b"\xff", b"\x00", b"<", b'"']


def _mutate(rng: random.Random, raw: bytes) -> bytes:
    """Mutations d'octets : remplacement, insertion de jetons XML, suppression, copie, nombre changé.
    Byte mutations: replacement, XML token insertion, deletion, copy, number changed."""
    data = bytearray(raw)
    for _ in range(rng.choice((1, 1, 2, 4))):
        at = rng.randrange(len(data)) if data else 0
        op = rng.randrange(5)
        if op == 0 and data:
            data[at] = rng.choice(b"<>&\"'/= \r\n\t:#x_0123456789AZaz\xc3\xa9")
        elif op == 1 or not data:
            data[at:at] = rng.choice(_MUTATION_TOKENS)
        elif op == 2:
            del data[at:at + rng.randint(1, 20)]
        elif op == 3:
            where = rng.randrange(len(data))
            data[where:where] = data[at:at + rng.randint(1, 60)]
        else:
            end = at
            while end < len(data) and data[end] in b"0123456789":
                end += 1
            number = rng.choice((0, 1, 60, 99, 65_536, 16_384, 4_294_967_296, rng.randint(0, 10**6)))
            data[at:end] = str(number).encode()
    return bytes(data)


def _guarded_compare(data: bytes) -> str:
    """`_compare` derrière la garde Python : calamine ne lit jamais une feuille que la garde écarterait (une
    cellule lointaine tuerait le processus). Sinon seul le noyau tourne : jamais d'autre exception que son refus.
    `_compare` behind the Python guard; otherwise only the core runs, raising nothing but its decline."""
    try:
        package = _package.open_package(data, "x.xlsx")
    except Exception:
        package = None
    try:
        largest = package.largest_extent() if package is not None else None
    except Exception:
        largest = None
    finally:
        if package is not None:
            package.close()
    if largest is not None and largest[0] * largest[1] <= _package.MAX_SHEET_AREA:
        return _compare(data)
    try:
        book = _CORE.XlsxBook(data)
    except _CORE.XlsxDecline:
        return "declined"
    for name, kind, _ in book.sheets:
        try:
            sheet = book.sheet(name) if kind == "worksheet" else None
            if sheet is not None and sheet.height * sheet.width < 1 << 20:
                sheet.rows()
        except (_CORE.XlsxDecline, OverflowError):
            continue
    return "unchecked"


def test_mutated_real_workbooks_read_exactly_as_calamine() -> None:
    """Classeurs xlsxwriter et openpyxl réels (textes Unicode retors, U+001C–U+001F, CRLF, dates 1900/1904,
    durées, formules en cache, fusions, feuilles masquées) mutés octet par octet ; seule la partie mutée change.
    Real xlsxwriter and openpyxl workbooks mutated byte by byte; only the mutated part changes."""
    import openpyxl

    texts = ["Élodie Ñúñez", "école", "\U0001F600 x", "a​b﻿c", "l1\r\nl2", "a\x1cb\x1f", "x" * 300,
             "[PERSONNE_1]", "a|b", "=SUM(A1)", "_x000D_"]

    def build(wb: Any) -> None:
        sheet = wb.add_worksheet("Données")
        dates = wb.add_format({"num_format": "dd/mm/yyyy hh:mm"})
        durations = wb.add_format({"num_format": "[h]:mm:ss"})
        for row, text in enumerate(texts):
            sheet.write_string(row, 0, text)
            sheet.write_number(row, 1, 40_000 + row * 7.25, dates)
            sheet.write_number(row, 2, row * 1.5 - 3, durations)
            sheet.write_formula(row, 3, f"=B{row + 1}*2", None, ["#N/A", 2.5, "calc", True][row % 4])
            sheet.write_boolean(row, 4, row % 2 == 0)
        sheet.merge_range("F1:G3", "fusion")
        wb.add_worksheet("Masquée").write("C5", "caché")
        wb.worksheets()[1].hide()

    workbook = openpyxl.Workbook()
    workbook.epoch = openpyxl.utils.datetime.CALENDAR_MAC_1904
    workbook.active.append(["Nom", dt.datetime(2021, 3, 4, 5, 6, 7), dt.time(1, 2, 3), dt.timedelta(hours=30), True])
    workbook.active.merge_cells("A3:B4")
    buffer = io.BytesIO()
    workbook.save(buffer)
    seeds = [xlsx_bytes(build), xlsx_bytes(build, {"date_1904": True}), buffer.getvalue()]
    for seed in seeds:
        assert _compare(seed) == "same"
    rng = random.Random(2024)
    outcomes: dict[str, int] = {}
    for _ in range(300):
        archive = zipfile.ZipFile(io.BytesIO(rng.choice(seeds)))
        members = [(info.filename, archive.read(info.filename)) for info in archive.infolist()]
        weights = [6 if name.startswith("xl/worksheets/") else 3 if "Strings" in name or "styles" in name else 1
                   for name, _ in members]
        index = rng.choices(range(len(members)), weights)[0]
        members[index] = (members[index][0], _mutate(rng, members[index][1]))
        outcome = _guarded_compare(zip_bytes(members))
        outcomes[outcome] = outcomes.get(outcome, 0) + 1
    assert outcomes.get("same", 0) + outcomes.get("partial", 0) >= 150, outcomes
