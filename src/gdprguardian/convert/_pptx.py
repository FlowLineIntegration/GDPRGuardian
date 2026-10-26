"""PPTX -> Markdown avec python-pptx / PPTX -> Markdown with python-pptx.

Une section par diapositive (titre, texte, tableaux GFM, graphiques en tableau, SmartArt, groupes),
notes selon `include_speaker_notes`, diapositives masquées selon `include_hidden`, commentaires selon
`include_comments` (sans auteur). Noms de formes, métadonnées et auteurs ne sont jamais émis.
One section per slide (title, text, GFM tables, charts as tables, SmartArt, groups), notes per
`include_speaker_notes`, hidden slides per `include_hidden`, comments per `include_comments` (no author).
Shape names, metadata and authors are never output.
"""

from __future__ import annotations

import io
from typing import Any

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.shapes.group import GroupShape
from pptx.shapes.picture import Movie, Picture

from ..documents import ConvertedDocument, DocumentOptions
from ._archive import SafeArchive, local_name, ooxml_metadata_count, parse_xml
from ._base import Report, Source, corrupt, too_large
from ._markdown import bullet_list, gfm_table, heading, image, inline, join_blocks, section

_P_NS = "{http://schemas.openxmlformats.org/presentationml/2006/main}"
_A_NS = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"
_DIAGRAM_IDS = "{http://schemas.openxmlformats.org/drawingml/2006/diagram}relIds"
_DIAGRAM_DATA = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}dm"
_BULLETED = frozenset({PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT})
_MAX_ROWS = 1000


class _Context:
    """État d'une conversion (budget de cellules, rapport) / per-conversion state (cell budget, report)."""

    __slots__ = ("cells", "options", "report", "slide")

    def __init__(self, options: DocumentOptions, report: Report) -> None:
        self.options = options
        self.report = report
        self.cells = 0
        self.slide: Any = None

    def spend(self, cells: int) -> None:
        self.cells += cells
        if self.cells > self.options.max_cells:
            raise too_large("cellules de tableau", "table cells", self.options.max_cells)


def convert(source: Source, options: DocumentOptions) -> ConvertedDocument:
    report = Report()
    try:
        presentation = Presentation(io.BytesIO(source.data))
        slides = list(presentation.slides)
    except Exception:  # python-pptx lève des types variés / python-pptx raises assorted types
        raise corrupt("pptx") from None
    if len(slides) > options.max_pages:
        raise too_large("diapositives", "slides", options.max_pages)
    context = _Context(options, report)
    blocks: list[str] = []
    for number, slide in enumerate(slides, 1):
        hidden = slide._element.get("show") in ("0", "false")
        if hidden and not options.include_hidden:
            report.drop("hidden_slides")
            continue
        context.slide = slide
        title_shape = slide.shapes.title
        title = ""
        if title_shape is not None and title_shape.has_text_frame:
            if options.include_hidden or not _is_hidden(title_shape):
                title = inline(title_shape.text_frame.text)
            else:
                report.drop("hidden_content")
        label = f"Diapositive {number}" + (" (masquée)" if hidden else "")
        blocks.append(heading(2, f"{label} : {title}" if title else label))
        skip = title_shape.shape_id if title_shape is not None else None
        blocks.extend(_shapes(slide.shapes, skip, context))
        blocks.append(_notes(slide, context))
        blocks.append(_comments(slide, context))
    report.drop("metadata", ooxml_metadata_count(source.archive or SafeArchive(source.data)))
    return report.document(join_blocks(blocks), "pptx", len(slides))


def _shapes(shapes: Any, skip: int | None, context: _Context) -> list[str]:
    blocks: list[str] = []
    for shape in shapes:
        if skip is not None and shape.shape_id == skip:
            continue
        if not context.options.include_hidden and _is_hidden(shape):
            context.report.drop("hidden_content")
            continue
        if isinstance(shape, GroupShape):
            blocks.extend(_shapes(shape.shapes, None, context))
        elif getattr(shape, "has_table", False):
            blocks.append(_table(shape.table, context))
        elif getattr(shape, "has_chart", False):
            blocks.append(_chart(shape.chart, context))
        elif isinstance(shape, (Picture, Movie)):
            context.report.drop("images")
            blocks.append(_picture(shape))
        elif shape.has_text_frame:
            blocks.append(_text_frame(shape.text_frame, _is_bulleted(shape)))
        elif getattr(shape._element, "graphicData_uri", None) == _DIAGRAM_URI:
            blocks.append(_smartart(shape, context))
    return blocks


def _properties(shape: Any) -> Any:
    """`p:cNvPr` de la forme (nom, description, masquage) / the shape's `p:cNvPr` (name, description, hidden)."""
    return shape._element.find(f"./*/{_P_NS}cNvPr")


def _is_hidden(shape: Any) -> bool:
    properties = _properties(shape)
    return properties is not None and properties.get("hidden") in ("1", "true")


def _is_bulleted(shape: Any) -> bool:
    try:
        return shape.is_placeholder and shape.placeholder_format.type in _BULLETED
    except (ValueError, AttributeError):
        return False


def _text_frame(frame: Any, bulleted: bool) -> str:
    lines = []
    for paragraph in frame.paragraphs:
        text = paragraph.text.replace("\v", " " if bulleted else "\n").strip()
        if text:
            lines.append(f"{'  ' * min(paragraph.level, 8)}- {text}" if bulleted else text)
    return "\n".join(lines)


def _table(table: Any, context: _Context) -> str:
    rows = list(table.rows)
    if len(rows) > _MAX_ROWS:
        context.report.warn(f"Tableau tronqué à {_MAX_ROWS} lignes / table truncated to {_MAX_ROWS} rows")
    cells = [[cell.text for cell in row.cells] for row in rows[:_MAX_ROWS]]
    context.spend(sum(map(len, cells)))
    return gfm_table(cells)


def _chart(chart: Any, context: _Context) -> str:
    """Données du graphique en tableau (catégories x séries) / chart data as a table (categories x series)."""
    try:
        title = inline(chart.chart_title.text_frame.text) if chart.has_title else ""
        plot = chart.plots[0]
        series = list(plot.series)
        values = [list(item.values) for item in series]
        categories = [str(category) for category in plot.categories]
        count = min(max([len(categories), *map(len, values)]), _MAX_ROWS)
        rows = [["Catégorie", *(inline(str(item.name)) for item in series)]]
        for index in range(count):
            label = categories[index] if index < len(categories) else str(index + 1)
            rows.append([label, *(_number(column[index]) if index < len(column) else "" for column in values)])
    except Exception:  # graphique incomplet ou données externes / incomplete chart or external data
        return "[GRAPHIQUE]"
    context.spend(sum(len(row) for row in rows))
    return join_blocks([f"Graphique : {title}" if title else "Graphique :", gfm_table(rows)])


def _number(value: Any) -> str:
    if value is None:
        return ""
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value)


def _picture(shape: Any) -> str:
    properties = _properties(shape)
    return image(properties.get("descr", "") if properties is not None else "")


def _smartart(shape: Any, context: _Context) -> str:
    ids = shape._element.find(f".//{_DIAGRAM_IDS}")
    try:
        part = context.slide.part.related_part(ids.get(_DIAGRAM_DATA))
        root = parse_xml(part.blob, "pptx")
    except (AttributeError, KeyError, TypeError):
        return ""
    texts = [inline("".join(t.text or "" for t in p.iter(f"{_A_NS}t"))) for p in root.iter(f"{_A_NS}p")]
    return bullet_list(texts)


def _notes(slide: Any, context: _Context) -> str:
    # Ne jamais lire `notes_slide` sans `has_notes_slide` : python-pptx en créerait une.
    # Never read `notes_slide` without `has_notes_slide`: python-pptx would create one.
    if not slide.has_notes_slide:
        return ""
    frame = slide.notes_slide.notes_text_frame
    text = _text_frame(frame, bulleted=False) if frame is not None else ""
    if not text:
        return ""
    if not context.options.include_speaker_notes:
        context.report.drop("speaker_notes")
        return ""
    return section("Notes", text, level=3)


def _comments(slide: Any, context: _Context) -> str:
    """Commentaires classiques et modernes, texte seul / legacy and modern comments, text only."""
    found: list[str] = []
    for rel in slide.part.rels.values():
        if rel.is_external or not rel.reltype.endswith("/comments"):
            continue
        root = parse_xml(rel.target_part.blob, "pptx")
        found.extend(inline(" ".join(el.text or "" for el in comment.iter() if local_name(el.tag) in ("text", "t")))
                     for comment in root.iter() if local_name(comment.tag) == "cm")
    if not context.options.include_comments:
        context.report.drop("comments", len(found))
        return ""
    return section("Commentaires", bullet_list(found), level=3)


__all__ = ["convert"]
